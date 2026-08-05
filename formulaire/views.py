"""
formulaire/views.py
Vues principales de l'application FCP / PADISAM
"""

from functools import wraps
import json

from django.contrib import messages
from django.contrib.auth import logout
from django.forms import inlineformset_factory
from django.http import JsonResponse, QueryDict, HttpResponse
from django.shortcuts import get_object_or_404, redirect, render
from django.urls import reverse
from django.utils import timezone
from django.utils.http import url_has_allowed_host_and_scheme
from django.db import transaction
from django.db.models import Sum, Count, Case, When, IntegerField, Value, Q
from django.core.paginator import Paginator
from django.db.models.functions import Cast
from django.contrib.auth.decorators import login_required


from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from .auth_forms import LoginForm
from .forms import (
    SousProjetForm,
    InfrastructureFormSet,
    EquipementFormSet,
    PromoteurFinalForm,
    IntrantFormSet,
    FonctionnementFormSet,
    ServiceFormSet,
    RealisationFormSet,
    EmpruntFormSet,
    ResultatPreselectionFormSet, 
    DecisionComiteSousProjetForm
)
from .models import (
    SousProjet,
    ResultatPreselection,
    Wilaya,
    Moughataa,
    Commune,
    Paysage,
    Village,
    Utilisateur,
    Infrastructure,
    Equipement,
    Intrant,
    Fonctionnement,
    Service,
    Activite,
    RealisationPassee,
    PassifEmprunt,
    HistoriqueSousProjet,
    HistoriqueModificationDossier,
)


# =========================================================
# OUTILS GÉNÉRAUX
# =========================================================

def get_current_user(request):
    """Récupère l'utilisateur connecté à partir de la session."""
    user_id = request.session.get('user_id')
    if not user_id:
        return None

    try:
        return Utilisateur.objects.select_related('wilaya').get(id=user_id, actif=True)
    except Utilisateur.DoesNotExist:
        return None


def is_agent_saisie(utilisateur):
    """Retourne True si l'utilisateur est un agent de saisie."""
    return bool(utilisateur and utilisateur.role == 'agent')
def is_superadmin(utilisateur):
    return bool(utilisateur and utilisateur.role == 'superadmin')


def is_admin(utilisateur):
    return bool(utilisateur and utilisateur.role == 'admin')


def is_admin_like(utilisateur):
    """
    Admin-like = admin + superadmin
    """
    return bool(utilisateur and utilisateur.role in ['admin', 'superadmin'])


# =========================================================
# RESTRICTIONS DU RÔLE CORRECTION PAR USERNAME
# =========================================================

CORRECTION_PAYSAGES_PAR_USERNAME = {
    'smaziz': [
        'CPB Extension',
        'Bababé',
    ],
}


def get_paysages_autorises_correction(utilisateur):
    """Retourne les paysages autorisés selon le username."""
    if not utilisateur or utilisateur.role != 'correction':
        return []

    username = (utilisateur.username or '').strip().lower()

    return CORRECTION_PAYSAGES_PAR_USERNAME.get(
        username,
        []
    )

def can_modify_project(utilisateur, sous_projet):
    """
    - admin / superadmin : tous les dossiers
    - correction : seulement les paysages autorisés pour son username
    - agent : seulement ses propres dossiers
    """
    if not utilisateur or not sous_projet:
        return False

    if utilisateur.role in ['admin', 'superadmin']:
        return True

    if utilisateur.role == 'correction':
        paysages_autorises = get_paysages_autorises_correction(
            utilisateur
        )

        if not paysages_autorises or not sous_projet.paysage:
            return False

        paysage_nom = (
            sous_projet.paysage.nom or ''
        ).strip().casefold()

        return any(
            paysage_nom == (nom or '').strip().casefold()
            for nom in paysages_autorises
        )

    if utilisateur.role == 'agent':
        return (
            (sous_projet.createur_username or '').strip().casefold()
            == (utilisateur.username or '').strip().casefold()
        )

    return False

def is_comite(utilisateur):
    """
    Accès comité = prescomite + superadmin
    """
    return bool(utilisateur and utilisateur.role in ['prescomite', 'superadmin'])


def can_delete_projects(utilisateur):
    return is_admin_like(utilisateur)


def can_manage_all_projects(utilisateur):
    return is_admin_like(utilisateur)


def can_validate_comite(utilisateur):
    return bool(utilisateur and utilisateur.role in ['prescomite', 'superadmin'])

def can_view_comite_list(utilisateur):
    return bool(utilisateur and utilisateur.role in [
        'admin', 'agent', 'superviseur', 'superadmin', 'prescomite'
    ])

def get_accessible_sous_projets(utilisateur):
    """
    - admin, superadmin, superviseur : tous les dossiers
    - correction : seulement les paysages autorisés pour son username
    - agent : seulement ses dossiers
    - prescomite : dossiers de sa wilaya
    """
    if not utilisateur:
        return SousProjet.objects.none()

    if utilisateur.role in [
        'admin',
        'superadmin',
        'superviseur',
    ]:
        return SousProjet.objects.all()

    if utilisateur.role == 'correction':
        paysages_autorises = get_paysages_autorises_correction(
            utilisateur
        )

        if not paysages_autorises:
            return SousProjet.objects.none()

        filtre_paysages = Q()

        for nom_paysage in paysages_autorises:
            nom_paysage = (nom_paysage or '').strip()

            if nom_paysage:
                filtre_paysages |= Q(
                    paysage__nom__iexact=nom_paysage
                )

        if not filtre_paysages:
            return SousProjet.objects.none()

        return SousProjet.objects.filter(
            filtre_paysages
        ).distinct()

    if utilisateur.role == 'agent':
        username = (utilisateur.username or '').strip()

        return SousProjet.objects.filter(
            createur_username__iexact=username
        )

    if (
        utilisateur.role == 'prescomite'
        and utilisateur.wilaya_id
    ):
        return SousProjet.objects.filter(
            wilaya_id=utilisateur.wilaya_id
        )

    return SousProjet.objects.none()

def get_current_sous_projet(request):
    """
    Récupère le sous-projet en cours depuis la session.
    Si le sous-projet n'est plus accessible, on nettoie la session.
    """
    sous_projet_id = request.session.get('current_sous_projet_id')
    utilisateur = get_current_user(request)

    if sous_projet_id and utilisateur:
        try:
            return get_accessible_sous_projets(utilisateur).get(id=sous_projet_id)
        except SousProjet.DoesNotExist:
            request.session.pop('current_sous_projet_id', None)
            return None

    return None


# =========================================================
# BROUILLON EN SESSION POUR LA CRÉATION MULTI-ÉTAPES
# =========================================================
# Objectif : pendant la création d'un nouveau sous-projet, aucune ligne n'est
# enregistrée dans la base de données avant la validation de la dernière étape.
# Les données validées étape par étape sont conservées dans la session.

NOUVEAU_SOUS_PROJET_DRAFT_KEY = 'nouveau_sous_projet_draft'


def _post_to_session_data(post_data):
    """Convertit un QueryDict POST en dict sérialisable pour la session."""
    return {key: post_data.getlist(key) for key in post_data.keys()}


def _session_data_to_querydict(data):
    """Reconstruit un QueryDict depuis les données conservées en session."""
    querydict = QueryDict('', mutable=True)
    for key, values in (data or {}).items():
        if isinstance(values, list):
            querydict.setlist(key, values)
        else:
            querydict.setlist(key, [values])
    return querydict


def _get_creation_draft(request):
    return request.session.get(NOUVEAU_SOUS_PROJET_DRAFT_KEY, {})


def _save_creation_step(request, step_name, post_data):
    draft = _get_creation_draft(request)
    draft[step_name] = _post_to_session_data(post_data)
    request.session[NOUVEAU_SOUS_PROJET_DRAFT_KEY] = draft
    request.session.modified = True


def _clear_creation_draft(request):
    request.session.pop(NOUVEAU_SOUS_PROJET_DRAFT_KEY, None)
    request.session.pop('current_sous_projet_id', None)
    request.session.modified = True


def _get_step_querydict(request, step_name):
    return _session_data_to_querydict(_get_creation_draft(request).get(step_name))


def _build_unsaved_sous_projet_from_step1(request, utilisateur):
    """
    Reconstruit un objet SousProjet non enregistré depuis l'étape 1.
    Sert uniquement à valider/afficher les formsets inline des étapes suivantes.
    """
    step1_data = _get_creation_draft(request).get('step1')
    if not step1_data:
        return None

    form = SousProjetForm(_session_data_to_querydict(step1_data), user=utilisateur)
    if not form.is_valid():
        return None

    sous_projet = form.save(commit=False)

    if is_agent_saisie(utilisateur):
        sous_projet.wilaya = utilisateur.wilaya

    if utilisateur:
        sous_projet.createur_username = utilisateur.username

    return sous_projet


def _financement_totals_from_formset(formset, include_subvention=True):
    """Calcule les totaux à partir des données du formset non encore sauvegardées."""
    totals = {
        'total_montant': 0,
        'total_contribution': 0,
        'total_autre': 0,
    }
    if include_subvention:
        totals['total_subvention'] = 0

    if formset.is_bound and formset.is_valid():
        for form in formset.forms:
            cleaned = getattr(form, 'cleaned_data', None) or {}
            if cleaned.get('DELETE'):
                continue
            if not cleaned.get('description') and cleaned.get('montant_total') in (None, ''):
                continue

            totals['total_montant'] += cleaned.get('montant_total') or 0
            totals['total_contribution'] += cleaned.get('contribution_promoteur') or 0
            totals['total_autre'] += cleaned.get('autre_financement') or 0
            if include_subvention:
                totals['total_subvention'] += cleaned.get('subvention_padisam') or 0

    return totals


def login_required(view_func):
    """Décorateur de protection des vues."""
    @wraps(view_func)
    def wrapper(request, *args, **kwargs):
        if not request.session.get('user_id'):
            return redirect('formulaire:login')
        return view_func(request, *args, **kwargs)
    return wrapper


# =========================================================
# OUTILS DE TOTALS FINANCEMENT
# =========================================================

def compute_financement_totals(queryset, include_subvention=True):
    """
    Calcule les totaux d'une section de financement.
    """
    if include_subvention:
        totals = queryset.aggregate(
            total_montant=Sum('montant_total'),
            total_subvention=Sum('subvention_padisam'),
            total_contribution=Sum('contribution_promoteur'),
            total_autre=Sum('autre_financement'),
        )
        return {
            'total_montant': totals['total_montant'] or 0,
            'total_subvention': totals['total_subvention'] or 0,
            'total_contribution': totals['total_contribution'] or 0,
            'total_autre': totals['total_autre'] or 0,
        }

    totals = queryset.aggregate(
        total_montant=Sum('montant_total'),
        total_contribution=Sum('contribution_promoteur'),
        total_autre=Sum('autre_financement'),
    )
    return {
        'total_montant': totals['total_montant'] or 0,
        'total_contribution': totals['total_contribution'] or 0,
        'total_autre': totals['total_autre'] or 0,
    }


# =========================================================
# FABRIQUES DE FORMSETS DYNAMIQUES
# =========================================================
# IMPORTANT :
# On garde les widgets/styles définis dans forms.py.
# On ne recrée PAS les formsets de financement à la main.
# On clone seulement la classe et on ajuste `extra`.

def clone_formset_with_extra(base_formset_class, extra_value):
    """
    Clone un formset existant en conservant ses widgets et options,
    puis remplace simplement la propriété extra.
    """
    CustomFormSet = type(
        f"Custom{base_formset_class.__name__}",
        (base_formset_class,),
        {'extra': extra_value}
    )
    return CustomFormSet


def get_activite_formset_class(extra=0):
    """
    Formset dynamique pour les activités.
    Ici on garde le formulaire custom ActiviteForm.
    """
    from .forms import ActiviteForm, BaseActiviteFormSet
    return inlineformset_factory(
        SousProjet,
        Activite,
        form=ActiviteForm,
        extra=extra,
        can_delete=True,
        formset=BaseActiviteFormSet
    )


# =========================================================
# AUTHENTIFICATION
# =========================================================

def login_view(request):
    """Page de connexion."""
    if request.session.get('user_id'):
        return redirect('formulaire:accueil')

    if request.method == 'POST':
        form = LoginForm(request.POST)

        if form.is_valid():
            username = form.cleaned_data['username']
            password = form.cleaned_data['password']

            try:
                utilisateur = Utilisateur.objects.get(username=username)

                if utilisateur.check_password(password) and utilisateur.actif:
                    request.session['user_id'] = utilisateur.id
                    request.session['user_name'] = f"{utilisateur.prenom} {utilisateur.nom}"
                    request.session['user_role'] = utilisateur.role
                    request.session['user_wilaya_id'] = utilisateur.wilaya_id
                    request.session['user_wilaya_nom'] = utilisateur.wilaya.nom if utilisateur.wilaya else ''

                    utilisateur.dernier_login = timezone.now()
                    utilisateur.save()

                    messages.success(request, f"Bienvenue {utilisateur.prenom} !")
                    return redirect('formulaire:accueil')
                else:
                    messages.error(request, "Mot de passe incorrect ou compte inactif.")

            except Utilisateur.DoesNotExist:
                messages.error(request, "Nom d'utilisateur inexistant.")
    else:
        form = LoginForm()

    return render(request, 'formulaire/login.html', {'form': form})


def logout_view(request):
    """Déconnexion."""
    logout(request)
    request.session.flush()
    messages.success(request, "Vous avez été déconnecté.")
    return redirect('formulaire:login')


# =========================================================
# ACCUEIL
# =========================================================

@login_required
def accueil(request):
    """Page d'accueil."""
    utilisateur = get_current_user(request)
    sous_projets = get_accessible_sous_projets(utilisateur)

    mode_label = ""
    wilaya_label = ""

    if utilisateur:
        if utilisateur.role == 'superadmin':
            mode_label = "Super administrateur"
        elif utilisateur.role == 'admin':
            mode_label = "Administrateur"
        elif utilisateur.role == 'agent':
            mode_label = "Saisie"
        elif utilisateur.role == 'prescomite':
            mode_label = "Président du comité de présélection"
        elif utilisateur.role == 'superviseur':
            mode_label = "Visiteur"
        elif utilisateur.role == 'correction':
            mode_label = "Correction des dossiers"
        elif utilisateur.role == 'consultant':
            mode_label = "Consultant"
        else:
            mode_label = utilisateur.role

        if utilisateur.role not in ['admin', 'superadmin'] and utilisateur.wilaya:
            wilaya_label = utilisateur.wilaya.nom

    context = {
        'user_name': request.session.get('user_name'),
        'user_role': request.session.get('user_role'),
        'user_wilaya_nom': request.session.get('user_wilaya_nom'),
        'mode_label': mode_label,
        'wilaya_label': wilaya_label,
        'total_projets': sous_projets.count(),
        'derniers_projets': sous_projets.order_by('-date_creation')[:5],
        'now': timezone.now(),
    }
    return render(request, 'formulaire/accueil.html', context)
# =========================================================
# ÉTAPE 1 : INFORMATIONS GÉNÉRALES + ACTIVITÉS
# =========================================================


@login_required
def nouveau_sous_projet(request):
    """
    Étape 1 : informations générales + activités.

    Nouvelle règle : aucune donnée n'est enregistrée en base à cette étape.
    Les données validées sont stockées temporairement dans la session.
    """
    utilisateur = get_current_user(request)

    if utilisateur and utilisateur.role in ['prescomite', 'superviseur', 'correction']:
        messages.error(request,
        "❌ Vous n'êtes pas autorisé à créer un nouveau dossier."
    )
        return redirect('formulaire:accueil')

    if is_agent_saisie(utilisateur) and not utilisateur.wilaya_id:
        messages.error(request, "Cet agent n'a pas de wilaya affectée. Impossible de créer un sous-projet.")
        return redirect('formulaire:accueil')

    if request.method == 'POST':
        form = SousProjetForm(request.POST, user=utilisateur)
        ActiviteDynamicFormSet = get_activite_formset_class(extra=0)
        activite_formset = ActiviteDynamicFormSet(
            request.POST,
            instance=SousProjet(),
            prefix='activite'
        )

        if form.is_valid() and activite_formset.is_valid():
            _save_creation_step(request, 'step1', request.POST)
            messages.success(request, "✅ Étape 1 validée. Les données seront enregistrées définitivement à la dernière étape.")
            return redirect('formulaire:financement_infrastructure')

        for field, errors in form.errors.items():
            for error in errors:
                messages.error(request, f"❌ {field} : {error}")

        if activite_formset.non_form_errors():
            for error in activite_formset.non_form_errors():
                messages.error(request, f"❌ Activités : {error}")

    else:
        draft_step1 = _get_creation_draft(request).get('step1')

        if draft_step1:
            step1_querydict = _session_data_to_querydict(draft_step1)
            form = SousProjetForm(step1_querydict, user=utilisateur)
            ActiviteDynamicFormSet = get_activite_formset_class(extra=0)
            activite_formset = ActiviteDynamicFormSet(
                step1_querydict,
                instance=SousProjet(),
                prefix='activite'
            )
        else:
            # Nouveau flux : on nettoie tout ancien ID temporaire issu de l'ancien comportement.
            request.session.pop('current_sous_projet_id', None)
            form = SousProjetForm(user=utilisateur)
            ActiviteDynamicFormSet = get_activite_formset_class(extra=1)
            activite_formset = ActiviteDynamicFormSet(
                instance=SousProjet(),
                prefix='activite'
            )

    return render(request, 'formulaire/nv_sous_projet.html', {
        'form': form,
        'activite_formset_data': activite_formset,
        'sous_projet': None,
        'user_wilaya_nom': request.session.get('user_wilaya_nom'),
    })


@login_required
def save_sous_projet(request):
    """Alias de sauvegarde étape 1."""
    return nouveau_sous_projet(request)


@login_required
def statistiques(request):
    """
    Page statistiques dynamique :
    - Par défaut : statistiques globales
    - Filtre exclusif par Wilaya OU par Paysage / ZOCA
    - Le total financement utilise la somme réelle des montant_total
    """
    utilisateur = get_current_user(request)

    sous_projets_base = get_accessible_sous_projets(utilisateur)

    # =====================================================
    # Choix du type de filtre : wilaya OU paysage
    # =====================================================
    filter_by = request.GET.get('filter_by') or 'wilaya'

    selected_wilaya_id = ''
    selected_paysage_id = ''

    if filter_by == 'wilaya':
        selected_wilaya_id = request.GET.get('wilaya') or ''
    elif filter_by == 'paysage':
        selected_paysage_id = request.GET.get('paysage') or ''
    else:
        filter_by = 'wilaya'

    selected_wilaya = None
    selected_paysage = None

    # =====================================================
    # Listes disponibles pour les filtres
    # =====================================================
    wilayas_disponibles = (
        sous_projets_base
        .exclude(wilaya__isnull=True)
        .values('wilaya_id', 'wilaya__nom')
        .distinct()
        .order_by('wilaya__nom')
    )

    paysages_disponibles = (
        sous_projets_base
        .exclude(paysage__isnull=True)
        .values('paysage_id', 'paysage__nom', 'wilaya__nom')
        .distinct()
        .order_by('wilaya__nom', 'paysage__nom')
    )

    # =====================================================
    # Application du filtre choisi
    # =====================================================
    sous_projets = sous_projets_base

    if filter_by == 'wilaya' and selected_wilaya_id:
        sous_projets = sous_projets.filter(wilaya_id=selected_wilaya_id)

        try:
            selected_wilaya = Wilaya.objects.get(id=selected_wilaya_id)
        except Wilaya.DoesNotExist:
            selected_wilaya = None

    if filter_by == 'paysage' and selected_paysage_id:
        sous_projets = sous_projets.filter(paysage_id=selected_paysage_id)

        try:
            selected_paysage = Paysage.objects.get(id=selected_paysage_id)
        except Paysage.DoesNotExist:
            selected_paysage = None

    total_projets = sous_projets.count()

    type_codes = ['AG', 'EL', 'SER', 'ENV']

    type_labels_long = {
        'AG': 'Agriculture',
        'EL': 'Élevage',
        'SER': 'Service',
        'ENV': 'Environnement',
    }

    # =====================================================
    # 1. Statistiques Wilaya
    # =====================================================
    stats_wilayas = (
        sous_projets
        .values('wilaya__nom')
        .annotate(total=Count('id'))
        .order_by('wilaya__nom')
    )

    # =====================================================
    # 2. Tableau Wilaya / Paysage / Type
    # =====================================================
    stats_paysages_types_raw = (
        sous_projets
        .values('wilaya__nom', 'paysage__nom', 'type_projet')
        .annotate(total=Count('id'))
        .order_by('wilaya__nom', 'paysage__nom', 'type_projet')
    )

    wilaya_data = {}

    for wilaya in stats_wilayas:
        wilaya_nom = wilaya['wilaya__nom'] or "Non renseignée"

        wilaya_data[wilaya_nom] = {
            'wilaya': wilaya_nom,
            'total': wilaya['total'],
            'paysages_dict': {},
            'totaux_types': {code: 0 for code in type_codes},
        }

    for item in stats_paysages_types_raw:
        wilaya_nom = item['wilaya__nom'] or "Non renseignée"
        paysage_nom = item['paysage__nom'] or "Non renseigné"
        type_code = item['type_projet'] or "NR"
        total = item['total']

        if wilaya_nom not in wilaya_data:
            wilaya_data[wilaya_nom] = {
                'wilaya': wilaya_nom,
                'total': 0,
                'paysages_dict': {},
                'totaux_types': {code: 0 for code in type_codes},
            }

        if paysage_nom not in wilaya_data[wilaya_nom]['paysages_dict']:
            wilaya_data[wilaya_nom]['paysages_dict'][paysage_nom] = {
                'paysage': paysage_nom,
                'types': {code: 0 for code in type_codes},
                'total': 0,
            }

        if type_code in type_codes:
            wilaya_data[wilaya_nom]['paysages_dict'][paysage_nom]['types'][type_code] += total
            wilaya_data[wilaya_nom]['totaux_types'][type_code] += total

        wilaya_data[wilaya_nom]['paysages_dict'][paysage_nom]['total'] += total

    stats_wilaya_paysages = []

    for wilaya_nom, data in wilaya_data.items():
        paysages = list(data['paysages_dict'].values())

        stats_wilaya_paysages.append({
            'wilaya': data['wilaya'],
            'total': data['total'],
            'paysages': paysages,
            'totaux_types': data['totaux_types'],
        })

    # =====================================================
    # 3. Pie chart : Répartition des projets par type
    # =====================================================
    stats_types_raw = (
        sous_projets
        .values('type_projet')
        .annotate(total=Count('id'))
        .order_by('type_projet')
    )

    stats_types = []
    chart_type_labels = []
    chart_type_data = []

    for item in stats_types_raw:
        code = item['type_projet'] or "NR"
        label = type_labels_long.get(code, code)
        total = item['total']

        pourcentage = round((total / total_projets) * 100, 2) if total_projets else 0

        stats_types.append({
            'type': label,
            'total': total,
            'pourcentage': pourcentage,
        })

        chart_type_labels.append(f"{label} : {total} projet(s) - {pourcentage}%")
        chart_type_data.append(total)

    # =====================================================
    # 4. Deuxième pie chart dynamique
    # =====================================================
    chart_second_labels = []
    chart_second_data = []

    if filter_by == 'wilaya' and selected_wilaya_id:
        second_chart_title = "Répartition des projets par paysage"

        stats_paysages_chart = (
            sous_projets
            .values('paysage__nom')
            .annotate(total=Count('id'))
            .order_by('paysage__nom')
        )

        for item in stats_paysages_chart:
            paysage_nom = item['paysage__nom'] or "Non renseigné"
            total = item['total']
            pourcentage = round((total / total_projets) * 100, 2) if total_projets else 0

            chart_second_labels.append(f"{paysage_nom} : {total} projet(s) - {pourcentage}%")
            chart_second_data.append(total)

    else:
        second_chart_title = "Répartition des projets par Wilaya"

        stats_wilayas_chart = (
            sous_projets
            .values('wilaya__nom')
            .annotate(total=Count('id'))
            .order_by('wilaya__nom')
        )

        for item in stats_wilayas_chart:
            wilaya_nom = item['wilaya__nom'] or "Non renseignée"
            total = item['total']
            pourcentage = round((total / total_projets) * 100, 2) if total_projets else 0

            chart_second_labels.append(f"{wilaya_nom} : {total} projet(s) - {pourcentage}%")
            chart_second_data.append(total)

    # =====================================================
    # 5. Statistiques globales des financements par Wilaya
    # =====================================================
    def total_relation(queryset, relation, champ):
        return queryset.aggregate(
            total=Sum(f'{relation}__{champ}')
        )['total'] or 0

    stats_financement_wilayas = []

    chart_financement_labels = []
    chart_financement_subvention = []
    chart_financement_contribution = []
    chart_financement_autre = []
    chart_financement_total = []

    wilayas_pour_financement = (
        sous_projets
        .exclude(wilaya__isnull=True)
        .values('wilaya_id', 'wilaya__nom')
        .distinct()
        .order_by('wilaya__nom')
    )

    for wilaya in wilayas_pour_financement:
        wilaya_id = wilaya['wilaya_id']
        wilaya_nom = wilaya['wilaya__nom'] or "Non renseignée"

        qs_wilaya = sous_projets.filter(wilaya_id=wilaya_id)

        total_montant = (
            total_relation(qs_wilaya, 'infrastructures', 'montant_total') +
            total_relation(qs_wilaya, 'equipements', 'montant_total') +
            total_relation(qs_wilaya, 'intrants', 'montant_total') +
            total_relation(qs_wilaya, 'fonctionnements', 'montant_total') +
            total_relation(qs_wilaya, 'services', 'montant_total')
        )

        total_subvention = (
            total_relation(qs_wilaya, 'infrastructures', 'subvention_padisam') +
            total_relation(qs_wilaya, 'equipements', 'subvention_padisam') +
            total_relation(qs_wilaya, 'intrants', 'subvention_padisam') +
            total_relation(qs_wilaya, 'services', 'subvention_padisam')
        )

        total_contribution = (
            total_relation(qs_wilaya, 'infrastructures', 'contribution_promoteur') +
            total_relation(qs_wilaya, 'equipements', 'contribution_promoteur') +
            total_relation(qs_wilaya, 'intrants', 'contribution_promoteur') +
            total_relation(qs_wilaya, 'fonctionnements', 'contribution_promoteur') +
            total_relation(qs_wilaya, 'services', 'contribution_promoteur')
        )

        total_autre = (
            total_relation(qs_wilaya, 'infrastructures', 'autre_financement') +
            total_relation(qs_wilaya, 'equipements', 'autre_financement') +
            total_relation(qs_wilaya, 'intrants', 'autre_financement') +
            total_relation(qs_wilaya, 'fonctionnements', 'autre_financement') +
            total_relation(qs_wilaya, 'services', 'autre_financement')
        )

        stats_financement_wilayas.append({
            'wilaya': wilaya_nom,
            'subvention': total_subvention,
            'contribution': total_contribution,
            'autre': total_autre,
            'total': total_montant,
        })

        chart_financement_labels.append(wilaya_nom)
        chart_financement_subvention.append(float(total_subvention))
        chart_financement_contribution.append(float(total_contribution))
        chart_financement_autre.append(float(total_autre))
        chart_financement_total.append(float(total_montant))

    total_financement_subvention = sum(item['subvention'] for item in stats_financement_wilayas)
    total_financement_contribution = sum(item['contribution'] for item in stats_financement_wilayas)
    total_financement_autre = sum(item['autre'] for item in stats_financement_wilayas)
    total_financement_general = sum(item['total'] for item in stats_financement_wilayas)

    # =====================================================
    # 6. Statistiques Agents de saisie
    # =====================================================
    stats_agents_raw = (
        sous_projets
        .values('createur_username')
        .annotate(total=Count('id'))
        .order_by('-total', 'createur_username')
    )

    stats_agents = []

    for item in stats_agents_raw:
        username = item['createur_username']

        nom_complet = "Non renseigné"
        wilaya_nom = "-"

        if username:
            try:
                user_obj = Utilisateur.objects.select_related('wilaya').get(username=username)

                nom_complet = f"{user_obj.prenom} {user_obj.nom}"

                if user_obj.wilaya:
                    wilaya_nom = user_obj.wilaya.nom

            except Utilisateur.DoesNotExist:
                nom_complet = "Utilisateur introuvable"

        stats_agents.append({
            'login': username or "-",
            'nom': nom_complet,
            'wilaya': wilaya_nom,
            'total': item['total'],
        })

    context = {
        'filter_by': filter_by,

        'total_projets': total_projets,

        'selected_wilaya_id': str(selected_wilaya_id),
        'selected_paysage_id': str(selected_paysage_id),

        'selected_wilaya': selected_wilaya,
        'selected_paysage': selected_paysage,

        'wilayas_disponibles': wilayas_disponibles,
        'paysages_disponibles': paysages_disponibles,

        'stats_wilaya_paysages': stats_wilaya_paysages,

        'stats_types': stats_types,

        'chart_type_labels': json.dumps(chart_type_labels),
        'chart_type_data': json.dumps(chart_type_data),

        'second_chart_title': second_chart_title,
        'chart_second_labels': json.dumps(chart_second_labels),
        'chart_second_data': json.dumps(chart_second_data),

        'stats_financement_wilayas': stats_financement_wilayas,
        'total_financement_subvention': total_financement_subvention,
        'total_financement_contribution': total_financement_contribution,
        'total_financement_autre': total_financement_autre,
        'total_financement_general': total_financement_general,

        'chart_financement_labels': json.dumps(chart_financement_labels),
        'chart_financement_subvention': json.dumps(chart_financement_subvention),
        'chart_financement_contribution': json.dumps(chart_financement_contribution),
        'chart_financement_autre': json.dumps(chart_financement_autre),
        'chart_financement_total': json.dumps(chart_financement_total),

        'stats_agents': stats_agents,

        'user_name': request.session.get('user_name'),
        'user_role': request.session.get('user_role'),
        'user_wilaya_nom': request.session.get('user_wilaya_nom'),
    }

    return render(request, 'formulaire/statistiques.html', context)
# =========================================================
# ÉTAPE 2 : INFRASTRUCTURES
# =========================================================

@login_required
def financement_infrastructure(request):
    utilisateur = get_current_user(request)
    sous_projet = _build_unsaved_sous_projet_from_step1(request, utilisateur)
    if not sous_projet:
        messages.error(request, "❌ Veuillez compléter l'étape 1 avant de continuer.")
        return redirect('formulaire:nouveau_sous_projet')

    if request.method == 'POST':
        DynamicInfrastructureFormSet = clone_formset_with_extra(InfrastructureFormSet, 0)
        formset = DynamicInfrastructureFormSet(request.POST, instance=sous_projet, prefix='infra')
        if formset.is_valid():
            _save_creation_step(request, 'infrastructures', request.POST)
            messages.success(request, "✅ Infrastructures validées temporairement.")
            return redirect('formulaire:financement_equipement')

        for form in formset.forms:
            for field, errors in form.errors.items():
                for error in errors:
                    messages.error(request, f"❌ Infrastructure - {field} : {error}")
    else:
        stored_data = _get_creation_draft(request).get('infrastructures')
        if stored_data:
            DynamicInfrastructureFormSet = clone_formset_with_extra(InfrastructureFormSet, 0)
            formset = DynamicInfrastructureFormSet(_session_data_to_querydict(stored_data), instance=sous_projet, prefix='infra')
        else:
            DynamicInfrastructureFormSet = clone_formset_with_extra(InfrastructureFormSet, 3)
            formset = DynamicInfrastructureFormSet(instance=sous_projet, prefix='infra')

    infrastructures_totaux = _financement_totals_from_formset(formset, include_subvention=True)

    return render(request, 'formulaire/financement_infrastructure.html', {
        'infrastructure_formset': formset,
        'sous_projet': sous_projet,
        'infrastructures_totaux': infrastructures_totaux,
    })


@login_required
def save_infrastructure(request):
    return financement_infrastructure(request)


# =========================================================
# ÉTAPE 3 : ÉQUIPEMENTS
# =========================================================

@login_required
def financement_equipement(request):
    utilisateur = get_current_user(request)
    sous_projet = _build_unsaved_sous_projet_from_step1(request, utilisateur)
    if not sous_projet:
        messages.error(request, "❌ Veuillez compléter l'étape 1 avant de continuer.")
        return redirect('formulaire:nouveau_sous_projet')

    if request.method == 'POST':
        DynamicEquipementFormSet = clone_formset_with_extra(EquipementFormSet, 0)
        formset = DynamicEquipementFormSet(request.POST, instance=sous_projet, prefix='equip')
        if formset.is_valid():
            _save_creation_step(request, 'equipements', request.POST)
            messages.success(request, "✅ Équipements validés temporairement.")
            return redirect('formulaire:financement_intrant')

        for form in formset.forms:
            for field, errors in form.errors.items():
                for error in errors:
                    messages.error(request, f"❌ Équipement - {field} : {error}")
    else:
        stored_data = _get_creation_draft(request).get('equipements')
        if stored_data:
            DynamicEquipementFormSet = clone_formset_with_extra(EquipementFormSet, 0)
            formset = DynamicEquipementFormSet(_session_data_to_querydict(stored_data), instance=sous_projet, prefix='equip')
        else:
            DynamicEquipementFormSet = clone_formset_with_extra(EquipementFormSet, 3)
            formset = DynamicEquipementFormSet(instance=sous_projet, prefix='equip')

    equipements_totaux = _financement_totals_from_formset(formset, include_subvention=True)

    return render(request, 'formulaire/financement_equipement.html', {
        'equipement_formset': formset,
        'sous_projet': sous_projet,
        'equipements_totaux': equipements_totaux,
    })


@login_required
def save_equipement(request):
    return financement_equipement(request)


# =========================================================
# ÉTAPE 4 : INTRANTS
# =========================================================

@login_required
def financement_intrant(request):
    utilisateur = get_current_user(request)
    sous_projet = _build_unsaved_sous_projet_from_step1(request, utilisateur)
    if not sous_projet:
        messages.error(request, "❌ Veuillez compléter l'étape 1 avant de continuer.")
        return redirect('formulaire:nouveau_sous_projet')

    if request.method == 'POST':
        DynamicIntrantFormSet = clone_formset_with_extra(IntrantFormSet, 0)
        formset = DynamicIntrantFormSet(request.POST, instance=sous_projet, prefix='intrant')
        if formset.is_valid():
            _save_creation_step(request, 'intrants', request.POST)
            messages.success(request, "✅ Intrants validés temporairement.")
            return redirect('formulaire:financement_fonctionnement')

        for form in formset.forms:
            for field, errors in form.errors.items():
                for error in errors:
                    messages.error(request, f"❌ Intrant - {field} : {error}")
    else:
        stored_data = _get_creation_draft(request).get('intrants')
        if stored_data:
            DynamicIntrantFormSet = clone_formset_with_extra(IntrantFormSet, 0)
            formset = DynamicIntrantFormSet(_session_data_to_querydict(stored_data), instance=sous_projet, prefix='intrant')
        else:
            DynamicIntrantFormSet = clone_formset_with_extra(IntrantFormSet, 3)
            formset = DynamicIntrantFormSet(instance=sous_projet, prefix='intrant')

    intrants_totaux = _financement_totals_from_formset(formset, include_subvention=True)

    return render(request, 'formulaire/financement_intrant.html', {
        'intrant_formset': formset,
        'sous_projet': sous_projet,
        'intrants_totaux': intrants_totaux,
    })


@login_required
def save_intrant(request):
    return financement_intrant(request)


# =========================================================
# ÉTAPE 5 : FONCTIONNEMENT
# =========================================================

@login_required
def financement_fonctionnement(request):
    utilisateur = get_current_user(request)
    sous_projet = _build_unsaved_sous_projet_from_step1(request, utilisateur)
    if not sous_projet:
        messages.error(request, "❌ Veuillez compléter l'étape 1 avant de continuer.")
        return redirect('formulaire:nouveau_sous_projet')

    if request.method == 'POST':
        DynamicFonctionnementFormSet = clone_formset_with_extra(FonctionnementFormSet, 0)
        formset = DynamicFonctionnementFormSet(request.POST, instance=sous_projet, prefix='fonc')
        if formset.is_valid():
            _save_creation_step(request, 'fonctionnements', request.POST)
            messages.success(request, "✅ Fonctionnement validé temporairement.")
            return redirect('formulaire:financement_services')

        for form in formset.forms:
            for field, errors in form.errors.items():
                for error in errors:
                    messages.error(request, f"❌ Fonctionnement - {field} : {error}")
    else:
        stored_data = _get_creation_draft(request).get('fonctionnements')
        if stored_data:
            DynamicFonctionnementFormSet = clone_formset_with_extra(FonctionnementFormSet, 0)
            formset = DynamicFonctionnementFormSet(_session_data_to_querydict(stored_data), instance=sous_projet, prefix='fonc')
        else:
            DynamicFonctionnementFormSet = clone_formset_with_extra(FonctionnementFormSet, 1)
            formset = DynamicFonctionnementFormSet(instance=sous_projet, prefix='fonc')

    fonctionnements_totaux = _financement_totals_from_formset(formset, include_subvention=False)

    return render(request, 'formulaire/financement_fonctionnement.html', {
        'fonctionnement_formset': formset,
        'sous_projet': sous_projet,
        'fonctionnements_totaux': fonctionnements_totaux,
    })


@login_required
def save_fonctionnement(request):
    return financement_fonctionnement(request)


# =========================================================
# ÉTAPE 6 : SERVICES
# =========================================================

@login_required
def financement_services(request):
    utilisateur = get_current_user(request)
    sous_projet = _build_unsaved_sous_projet_from_step1(request, utilisateur)
    if not sous_projet:
        messages.error(request, "❌ Veuillez compléter l'étape 1 avant de continuer.")
        return redirect('formulaire:nouveau_sous_projet')

    if request.method == 'POST':
        DynamicServiceFormSet = clone_formset_with_extra(ServiceFormSet, 0)
        formset = DynamicServiceFormSet(request.POST, instance=sous_projet, prefix='serv')
        if formset.is_valid():
            _save_creation_step(request, 'services', request.POST)
            messages.success(request, "✅ Services validés temporairement.")
            return redirect('formulaire:realisation_passif')

        for form in formset.forms:
            for field, errors in form.errors.items():
                for error in errors:
                    messages.error(request, f"❌ Service - {field} : {error}")
    else:
        stored_data = _get_creation_draft(request).get('services')
        if stored_data:
            DynamicServiceFormSet = clone_formset_with_extra(ServiceFormSet, 0)
            formset = DynamicServiceFormSet(_session_data_to_querydict(stored_data), instance=sous_projet, prefix='serv')
        else:
            DynamicServiceFormSet = clone_formset_with_extra(ServiceFormSet, 1)
            formset = DynamicServiceFormSet(instance=sous_projet, prefix='serv')

    services_totaux = _financement_totals_from_formset(formset, include_subvention=True)

    return render(request, 'formulaire/financement_services.html', {
        'service_formset': formset,
        'sous_projet': sous_projet,
        'services_totaux': services_totaux,
    })


@login_required
def save_service(request):
    return financement_services(request)


# =========================================================
# ÉTAPE 7 : RÉALISATIONS PASSÉES ET EMPRUNTS
# =========================================================

@login_required
def realisation_passif(request):
    """
    Dernière étape.
    C'est uniquement ici que le sous-projet et toutes ses tables liées sont créés en base.
    Si l'utilisateur abandonne avant cette étape, rien n'est enregistré.
    """
    utilisateur = get_current_user(request)
    sous_projet = _build_unsaved_sous_projet_from_step1(request, utilisateur)
    if not sous_projet:
        messages.error(request, "❌ Veuillez compléter l'étape 1 avant de continuer.")
        return redirect('formulaire:nouveau_sous_projet')

    draft = _get_creation_draft(request)
    required_steps = ['step1', 'infrastructures', 'equipements', 'intrants', 'fonctionnements', 'services']
    missing_steps = [step for step in required_steps if step not in draft]

    if missing_steps:
        messages.error(request, "❌ Veuillez compléter toutes les étapes avant l'enregistrement final.")
        return redirect('formulaire:nouveau_sous_projet')

    if request.method == 'POST':
        promoteur_form = PromoteurFinalForm(request.POST, instance=sous_projet)
        realisation_formset = RealisationFormSet(request.POST, prefix='real')
        emprunt_formset = EmpruntFormSet(request.POST, prefix='emprunt')

        # On revalide toutes les étapes conservées en session juste avant la création réelle.
        step1_qd = _get_step_querydict(request, 'step1')
        step1_form = SousProjetForm(step1_qd, user=utilisateur)
        step1_activite_formset = get_activite_formset_class(extra=0)(step1_qd, instance=SousProjet(), prefix='activite')

        infrastructure_formset = clone_formset_with_extra(InfrastructureFormSet, 0)(
            _get_step_querydict(request, 'infrastructures'), instance=sous_projet, prefix='infra'
        )
        equipement_formset = clone_formset_with_extra(EquipementFormSet, 0)(
            _get_step_querydict(request, 'equipements'), instance=sous_projet, prefix='equip'
        )
        intrant_formset = clone_formset_with_extra(IntrantFormSet, 0)(
            _get_step_querydict(request, 'intrants'), instance=sous_projet, prefix='intrant'
        )
        fonctionnement_formset = clone_formset_with_extra(FonctionnementFormSet, 0)(
            _get_step_querydict(request, 'fonctionnements'), instance=sous_projet, prefix='fonc'
        )
        service_formset = clone_formset_with_extra(ServiceFormSet, 0)(
            _get_step_querydict(request, 'services'), instance=sous_projet, prefix='serv'
        )

        all_valid = all([
            step1_form.is_valid(),
            step1_activite_formset.is_valid(),
            infrastructure_formset.is_valid(),
            equipement_formset.is_valid(),
            intrant_formset.is_valid(),
            fonctionnement_formset.is_valid(),
            service_formset.is_valid(),
            promoteur_form.is_valid(),
            realisation_formset.is_valid(),
            emprunt_formset.is_valid(),
        ])

        if all_valid:
            try:
                with transaction.atomic():
                    sous_projet = step1_form.save(commit=False)

                    if is_agent_saisie(utilisateur):
                        sous_projet.wilaya = utilisateur.wilaya

                    if utilisateur:
                        sous_projet.createur_username = utilisateur.username

                    # Ajoute les champs de la dernière étape sur le même objet.
                    promoteur_final = PromoteurFinalForm(request.POST, instance=sous_projet).save(commit=False)
                    sous_projet = promoteur_final
                    sous_projet.save()

                    step1_activite_formset.instance = sous_projet
                    step1_activite_formset.save()

                    for formset in [
                        infrastructure_formset,
                        equipement_formset,
                        intrant_formset,
                        fonctionnement_formset,
                        service_formset,
                    ]:
                        formset.instance = sous_projet
                        formset.save()

                    for real_form in realisation_formset:
                        cleaned_data = getattr(real_form, 'cleaned_data', None)
                        if not cleaned_data:
                            continue

                        has_data = any([
                            cleaned_data.get('produit'),
                            cleaned_data.get('annee_1'), cleaned_data.get('volume_annee_1'), cleaned_data.get('ventes_usd_annee_1'), cleaned_data.get('prix_vente_mru_annee_1'),
                            cleaned_data.get('annee_2'), cleaned_data.get('volume_annee_2'), cleaned_data.get('ventes_usd_annee_2'), cleaned_data.get('prix_vente_mru_annee_2'),
                            cleaned_data.get('annee_3'), cleaned_data.get('volume_annee_3'), cleaned_data.get('ventes_usd_annee_3'), cleaned_data.get('prix_vente_mru_annee_3'),
                        ])

                        if not has_data:
                            continue

                        RealisationPassee.objects.create(
                            sous_projet=sous_projet,
                            produit=cleaned_data.get('produit'),
                            annee_1=cleaned_data.get('annee_1'),
                            volume_annee_1=cleaned_data.get('volume_annee_1'),
                            ventes_usd_annee_1=cleaned_data.get('ventes_usd_annee_1'),
                            prix_vente_mru_annee_1=cleaned_data.get('prix_vente_mru_annee_1'),
                            annee_2=cleaned_data.get('annee_2'),
                            volume_annee_2=cleaned_data.get('volume_annee_2'),
                            ventes_usd_annee_2=cleaned_data.get('ventes_usd_annee_2'),
                            prix_vente_mru_annee_2=cleaned_data.get('prix_vente_mru_annee_2'),
                            annee_3=cleaned_data.get('annee_3'),
                            volume_annee_3=cleaned_data.get('volume_annee_3'),
                            ventes_usd_annee_3=cleaned_data.get('ventes_usd_annee_3'),
                            prix_vente_mru_annee_3=cleaned_data.get('prix_vente_mru_annee_3'),
                        )

                    for emp_form in emprunt_formset:
                        cleaned_data = getattr(emp_form, 'cleaned_data', None)
                        if not cleaned_data:
                            continue

                        has_data = any([
                            cleaned_data.get('annee'),
                            cleaned_data.get('institution_financiere'),
                            cleaned_data.get('montant_emprunte'),
                            cleaned_data.get('montant_rembourse'),
                        ])

                        if not has_data:
                            continue

                        PassifEmprunt.objects.create(
                            sous_projet=sous_projet,
                            annee=cleaned_data.get('annee'),
                            institution_financiere=cleaned_data.get('institution_financiere'),
                            montant_emprunte=cleaned_data.get('montant_emprunte'),
                            montant_rembourse=cleaned_data.get('montant_rembourse'),
                        )

                    enregistrer_historique_modification(
                        sous_projet=sous_projet,
                        utilisateur=utilisateur,
                        section='dossier_complet',
                        action='creation'
                    )

                _clear_creation_draft(request)
                messages.success(request, "✅ Sous-projet créé avec succès !")
                return redirect('formulaire:liste_sous_projets')

            except Exception as exc:
                messages.error(request, f"❌ Enregistrement annulé. Aucune donnée n'a été sauvegardée. Détail : {exc}")

        else:
            messages.error(request, "❌ Certaines étapes contiennent des erreurs. Aucune donnée n'a été sauvegardée.")

            for field, errors in promoteur_form.errors.items():
                for error in errors:
                    messages.error(request, f"❌ Promoteur - {field} : {error}")

            for form in realisation_formset.forms:
                for field, errors in form.errors.items():
                    for error in errors:
                        messages.error(request, f"❌ Réalisation - {field} : {error}")

            for form in emprunt_formset.forms:
                for field, errors in form.errors.items():
                    for error in errors:
                        messages.error(request, f"❌ Emprunt - {field} : {error}")

    else:
        stored_final = draft.get('final')
        if stored_final:
            final_qd = _session_data_to_querydict(stored_final)
            promoteur_form = PromoteurFinalForm(final_qd, instance=sous_projet)
            realisation_formset = RealisationFormSet(final_qd, prefix='real')
            emprunt_formset = EmpruntFormSet(final_qd, prefix='emprunt')
        else:
            promoteur_form = PromoteurFinalForm(instance=sous_projet)
            realisation_formset = RealisationFormSet(prefix='real')
            emprunt_formset = EmpruntFormSet(prefix='emprunt')

    return render(request, 'formulaire/realisation_passif.html', {
        'promoteur_form': promoteur_form,
        'realisation_formset': realisation_formset,
        'emprunt_formset': emprunt_formset,
        'sous_projet': sous_projet,
    })


@login_required
def save_realisation_passif(request):
    return realisation_passif(request)


# =========================================================
# LISTE / DÉTAIL / SUPPRESSION
# =========================================================

@login_required
def liste_sous_projets(request):
    """
    Liste des sous-projets accessibles.

    Par défaut :
    - affiche tous les sous-projets accessibles.

    Filtres disponibles :
    - par Wilaya
    - par Paysage / ZOCA
    - par Numéro réception formulaire
    """
    utilisateur = get_current_user(request)

    sous_projets_base = get_accessible_sous_projets(utilisateur)

    filter_by = request.GET.get('filter_by') or 'wilaya'

    selected_wilaya_id = ''
    selected_paysage_id = ''
    numero_reception = ''

    if filter_by == 'wilaya':
        selected_wilaya_id = request.GET.get('wilaya') or ''
    elif filter_by == 'paysage':
        selected_paysage_id = request.GET.get('paysage') or ''
    elif filter_by == 'numero':
        numero_reception = request.GET.get('numero_reception') or ''
    else:
        filter_by = 'wilaya'

    selected_wilaya = None
    selected_paysage = None

    wilayas_disponibles = (
        sous_projets_base
        .exclude(wilaya__isnull=True)
        .values('wilaya_id', 'wilaya__nom')
        .distinct()
        .order_by('wilaya__nom')
    )

    paysages_disponibles = (
        sous_projets_base
        .exclude(paysage__isnull=True)
        .values('paysage_id', 'paysage__nom', 'wilaya__nom')
        .distinct()
        .order_by('wilaya__nom', 'paysage__nom')
    )

    sous_projets = sous_projets_base

    if filter_by == 'wilaya' and selected_wilaya_id:
        sous_projets = sous_projets.filter(wilaya_id=selected_wilaya_id)

        try:
            selected_wilaya = Wilaya.objects.get(id=selected_wilaya_id)
        except Wilaya.DoesNotExist:
            selected_wilaya = None

    if filter_by == 'paysage' and selected_paysage_id:
        sous_projets = sous_projets.filter(paysage_id=selected_paysage_id)

        try:
            selected_paysage = Paysage.objects.get(id=selected_paysage_id)
        except Paysage.DoesNotExist:
            selected_paysage = None

    if filter_by == 'numero' and numero_reception:
        sous_projets = sous_projets.filter(
            numero_reception_formulaire__icontains=numero_reception
        )

    sous_projets = (
    sous_projets
    .select_related(
        'wilaya',
        'moughataa',
        'commune',
        'paysage'
    )
    .annotate(
        numero_reception_numerique=Case(
            When(
                numero_reception_formulaire__regex=r'^[0-9]+$',
                then=Cast(
                    'numero_reception_formulaire',
                    IntegerField()
                )
            ),
            default=Value(None),
            output_field=IntegerField()
        )
    )
    .order_by(
        'numero_reception_numerique',
        'numero_reception_formulaire',
        'id'
    )
)

    return render(request, 'formulaire/liste_sous_projets.html', {
        'filter_by': filter_by,

        'sous_projets': sous_projets,

        'wilayas_disponibles': wilayas_disponibles,
        'paysages_disponibles': paysages_disponibles,

        'selected_wilaya_id': str(selected_wilaya_id),
        'selected_paysage_id': str(selected_paysage_id),
        'numero_reception': numero_reception,

        'selected_wilaya': selected_wilaya,
        'selected_paysage': selected_paysage,

        'user_name': request.session.get('user_name'),
        'user_role': request.session.get('user_role'),
        'user_wilaya_nom': request.session.get('user_wilaya_nom'),
        'is_admin': is_admin_like(utilisateur),
        'is_superadmin': is_superadmin(utilisateur),
        'can_edit_all': bool(
            utilisateur
            and utilisateur.role in ['admin', 'superadmin']
        ),
        'is_correction': bool(
            utilisateur
            and utilisateur.role == 'correction'
        ),
    })

@login_required
def detail_sous_projet(request, pk):
    """Détail d'un sous-projet."""
    utilisateur = get_current_user(request)

    sous_projet = get_object_or_404(
        get_accessible_sous_projets(utilisateur),
        pk=pk
    )

    liste_retour = request.GET.get('next')

    if not liste_retour or not url_has_allowed_host_and_scheme(
        url=liste_retour,
        allowed_hosts={request.get_host()},
        require_https=request.is_secure(),
    ):
        liste_retour = reverse(
            'formulaire:liste_sous_projets'
        )

    infrastructures = list(sous_projet.infrastructures.all())
    equipements = list(sous_projet.equipements.all())
    intrants = list(sous_projet.intrants.all())
    fonctionnements = list(sous_projet.fonctionnements.all())
    services = list(sous_projet.services.all())

    total_infrastructures = sous_projet.infrastructures.aggregate(total=Sum('montant_total'))['total'] or 0
    total_equipements = sous_projet.equipements.aggregate(total=Sum('montant_total'))['total'] or 0
    total_intrants = sous_projet.intrants.aggregate(total=Sum('montant_total'))['total'] or 0
    total_fonctionnements = sous_projet.fonctionnements.aggregate(total=Sum('montant_total'))['total'] or 0
    total_services = sous_projet.services.aggregate(total=Sum('montant_total'))['total'] or 0

    # Totaux détaillés par source de financement
    infra_subvention = sous_projet.infrastructures.aggregate(total=Sum('subvention_padisam'))['total'] or 0
    infra_contribution = sous_projet.infrastructures.aggregate(total=Sum('contribution_promoteur'))['total'] or 0
    infra_autre = sous_projet.infrastructures.aggregate(total=Sum('autre_financement'))['total'] or 0

    equip_subvention = sous_projet.equipements.aggregate(total=Sum('subvention_padisam'))['total'] or 0
    equip_contribution = sous_projet.equipements.aggregate(total=Sum('contribution_promoteur'))['total'] or 0
    equip_autre = sous_projet.equipements.aggregate(total=Sum('autre_financement'))['total'] or 0

    intrant_subvention = sous_projet.intrants.aggregate(total=Sum('subvention_padisam'))['total'] or 0
    intrant_contribution = sous_projet.intrants.aggregate(total=Sum('contribution_promoteur'))['total'] or 0
    intrant_autre = sous_projet.intrants.aggregate(total=Sum('autre_financement'))['total'] or 0

    fonct_contribution = sous_projet.fonctionnements.aggregate(total=Sum('contribution_promoteur'))['total'] or 0
    fonct_autre = sous_projet.fonctionnements.aggregate(total=Sum('autre_financement'))['total'] or 0

    service_subvention = sous_projet.services.aggregate(total=Sum('subvention_padisam'))['total'] or 0
    service_contribution = sous_projet.services.aggregate(total=Sum('contribution_promoteur'))['total'] or 0
    service_autre = sous_projet.services.aggregate(total=Sum('autre_financement'))['total'] or 0
     # =====================================================
    # Tableau récapitulatif des financements
    # =====================================================

        
    recap_financement = [
        {
            'libelle': 'Infrastructures',
            'subvention': infra_subvention,
            'contribution': infra_contribution,
            'autre': infra_autre,
            'total': total_infrastructures,
        },
        {
            'libelle': 'Équipements',
            'subvention': equip_subvention,
            'contribution': equip_contribution,
            'autre': equip_autre,
            'total': total_equipements,
        },
        {
            'libelle': 'Intrants',
            'subvention': intrant_subvention,
            'contribution': intrant_contribution,
            'autre': intrant_autre,
            'total': total_intrants,
        },
        {
            'libelle': 'Fonctionnement',
            'subvention': 0,
            'contribution': fonct_contribution,
            'autre': fonct_autre,
            'total': total_fonctionnements,
        },
        {
            'libelle': 'Services',
            'subvention': service_subvention,
            'contribution': service_contribution,
            'autre': service_autre,
            'total': total_services,
        },
    ]

    recap_total_subvention = (
        infra_subvention +
        equip_subvention +
        intrant_subvention +
        service_subvention
    )

    recap_total_contribution = (
        fonct_contribution +
        infra_contribution +
        equip_contribution +
        service_contribution +
        intrant_contribution
    )

    recap_total_autre = (
        fonct_autre +
        infra_autre +
        equip_autre +
        service_autre +
        intrant_autre
    )

    recap_total_general = (
        total_fonctionnements +
        total_infrastructures +
        total_equipements +
        total_services +
        total_intrants
    )

    grand_total = (
        total_infrastructures +
        total_equipements +
        total_intrants +
        total_fonctionnements +
        total_services
    )

    total_ventes = 0
    for real in sous_projet.realisations.all():
        total_ventes += (real.ventes_usd_annee_1 or 0)
        total_ventes += (real.ventes_usd_annee_2 or 0)
        total_ventes += (real.ventes_usd_annee_3 or 0)

    total_emprunte = sum(emp.montant_emprunte or 0 for emp in sous_projet.emprunts.all())
    total_rembourse = sum(emp.montant_rembourse or 0 for emp in sous_projet.emprunts.all())

    context = {
        'sous_projet': sous_projet,
        'can_edit': can_modify_project(utilisateur, sous_projet),
        'user_name': request.session.get('user_name'),
        'user_role': request.session.get('user_role'),
        'user_wilaya_nom': request.session.get('user_wilaya_nom'),
        'liste_retour': liste_retour,

        'infrastructures': infrastructures,
        'equipements': equipements,
        'intrants': intrants,
        'fonctionnements': fonctionnements,
        'services': services,

        'total_infrastructures': total_infrastructures,
        'total_equipements': total_equipements,
        'total_intrants': total_intrants,
        'total_fonctionnements': total_fonctionnements,
        'total_services': total_services,

        'infra_subvention': infra_subvention,
        'infra_contribution': infra_contribution,
        'infra_autre': infra_autre,

        'equip_subvention': equip_subvention,
        'equip_contribution': equip_contribution,
        'equip_autre': equip_autre,

        'intrant_subvention': intrant_subvention,
        'intrant_contribution': intrant_contribution,
        'intrant_autre': intrant_autre,

        'fonct_contribution': fonct_contribution,
        'fonct_autre': fonct_autre,

        'service_subvention': service_subvention,
        'service_contribution': service_contribution,
        'service_autre': service_autre,

        'service_subvention': service_subvention,
        'service_contribution': service_contribution,
        'service_autre': service_autre,

        'recap_financement': recap_financement,
        'recap_total_subvention': recap_total_subvention,
        'recap_total_contribution': recap_total_contribution,
        'recap_total_autre': recap_total_autre,
        'recap_total_general': recap_total_general,

        'grand_total': grand_total,

        'grand_total': grand_total,
        'total_ventes': total_ventes,
        'total_emprunte': total_emprunte,
        'total_rembourse': total_rembourse,
    }

    return render(request, 'formulaire/detail_sous_projet.html', context)
@login_required
def supprimer_sous_projet(request, pk):
    """Supprime un sous-projet accessible."""
    utilisateur = get_current_user(request)

    if not can_delete_projects(utilisateur):
        messages.error(request, "❌ Seuls les administrateurs peuvent supprimer un sous-projet.")
        return redirect('formulaire:liste_sous_projets')

    sous_projet = get_object_or_404(get_accessible_sous_projets(utilisateur), pk=pk)

    if request.method == 'POST':
        intitule = sous_projet.intitule_sous_projet

        try:
            with transaction.atomic():
                enregistrer_historique_modification(
                    sous_projet=sous_projet,
                    utilisateur=utilisateur,
                    section='dossier_complet',
                    action='suppression'
                )
                sous_projet.delete()

            messages.success(
                request,
                f'✅ Sous-projet "{intitule}" supprimé avec succès !'
            )
            return redirect('formulaire:liste_sous_projets')

        except Exception as exc:
            messages.error(
                request,
                "❌ La suppression a été annulée : "
                f"{exc}"
            )
            return redirect(
                'formulaire:detail_sous_projet',
                pk=sous_projet.pk
            )

    return redirect('formulaire:liste_sous_projets')
# =========================================================
# API AJAX
# =========================================================

def get_moughataas(request):
    """Retourne les moughataas d'une wilaya."""
    wilaya_id = request.GET.get('wilaya_id')
    user_role = request.session.get('user_role')
    user_wilaya_id = request.session.get('user_wilaya_id')

    if user_role == 'agent' and user_wilaya_id and str(wilaya_id) != str(user_wilaya_id):
        return JsonResponse([], safe=False)

    if wilaya_id:
        moughataas = Moughataa.objects.filter(wilaya_id=wilaya_id).order_by('nom')
        data = [{'id': m.id, 'nom': m.nom} for m in moughataas]
        return JsonResponse(data, safe=False)

    return JsonResponse([], safe=False)


def get_communes(request):
    """Retourne les communes d'une moughataa."""
    moughataa_id = request.GET.get('moughataa_id')
    user_role = request.session.get('user_role')
    user_wilaya_id = request.session.get('user_wilaya_id')

    if moughataa_id and moughataa_id.isdigit():
        try:
            communes = Commune.objects.filter(moughataa_id=moughataa_id)

            if user_role == 'agent' and user_wilaya_id:
                communes = communes.filter(moughataa__wilaya_id=user_wilaya_id)

            data = [{'id': c.id, 'nom': c.nom} for c in communes.order_by('nom')]
            return JsonResponse(data, safe=False)
        except Exception:
            return JsonResponse([], safe=False)

    return JsonResponse([], safe=False)


def get_paysages(request):
    """Retourne les paysages d'une commune."""
    commune_id = request.GET.get('commune_id')
    user_role = request.session.get('user_role')
    user_wilaya_id = request.session.get('user_wilaya_id')

    if commune_id and commune_id.isdigit():
        try:
            paysages = Paysage.objects.filter(commune_id=commune_id)

            if user_role == 'agent' and user_wilaya_id:
                paysages = paysages.filter(commune__moughataa__wilaya_id=user_wilaya_id)

            data = [{'id': p.id, 'nom': p.nom} for p in paysages.order_by('nom')]
            return JsonResponse(data, safe=False)
        except Exception:
            return JsonResponse([], safe=False)

    return JsonResponse([], safe=False)


def get_villages(request):
    """Retourne les villages d'un paysage."""
    paysage_id = request.GET.get('paysage_id')

    if paysage_id and paysage_id.isdigit():
        try:
            villages = Village.objects.filter(paysage_id=paysage_id).order_by('nom')
            data = [{'nom': v.nom} for v in villages]
            return JsonResponse(data, safe=False)
        except Exception:
            return JsonResponse([], safe=False)

    return JsonResponse([], safe=False)


def evaluer_preselection(sp):
    """
    Présélection automatique simple.
    Champs obligatoires éliminatoires :
    - Intitulé du sous-projet
    - Objectif du sous-projet
    """

    motifs_manquants = []

    if not sp.intitule_sous_projet:
        motifs_manquants.append("Intitulé du sous-projet non renseigné")

    if not sp.objectif_sous_projet:
        motifs_manquants.append("Objectif du sous-projet non renseigné")

    if motifs_manquants:
        decision = "Rejeté"
        motif = "Rejet automatique : " + "; ".join(motifs_manquants) + "."
        badge_class = "badge-no"
    else:
        decision = "À examiner"
        motif = "Le dossier contient l’intitulé et l’objectif du sous-projet. Il peut être examiné par le comité."
        badge_class = "badge-review"

    return {
        "numero_reception": sp.numero_reception_formulaire or "Non renseigné",
        "intitule": sp.intitule_sous_projet or "Non renseigné",
        "decision": decision,
        "motif": motif,
        "badge_class": badge_class,
        "sous_projet": sp,
    }

def appliquer_preselection_automatique(sp):
    motifs_manquants = []

    if not sp.intitule_sous_projet:
        motifs_manquants.append("Intitulé du sous-projet non renseigné")

    if not sp.objectif_sous_projet:
        motifs_manquants.append("Objectif du sous-projet non renseigné")

    if motifs_manquants:
        return {
            'decision': 'Rejet automatique',
            'motif': "Informations manquantes : " + "; ".join(motifs_manquants) + ".",
            'badge_class': 'badge-no',
            'bloque_comite': False,
        }

    return {
        'decision': 'À examiner',
        'motif': "Le dossier contient l’intitulé et l’objectif du sous-projet. Il peut être examiné par le comité.",
        'badge_class': 'badge-review',
        'bloque_comite': False,
    }



def normaliser_statut_preselection(statut):
    """
    Normalise les valeurs utilisées dans l'historique.
    Les valeurs vides et le statut technique « etude » sont considérés
    comme « a_examiner ».
    """
    statut = (statut or '').strip()

    if statut in ['', 'etude', 'a_examiner']:
        return 'a_examiner'

    return statut


def enregistrer_historique_statut(
    sous_projet,
    utilisateur,
    ancien_statut,
    nouveau_statut
):
    """
    Enregistre uniquement un changement réel de statut du dossier.

    Les informations du dossier sont copiées au moment de l'action :
    numéro de réception, intitulé, paysage, ancien statut,
    nouveau statut, utilisateur, date et heure.
    """
    if not sous_projet or not utilisateur:
        return None

    ancien_statut = normaliser_statut_preselection(
        ancien_statut
    )
    nouveau_statut = normaliser_statut_preselection(
        nouveau_statut
    )

    # Ne pas créer une ligne si le statut n'a pas changé.
    if ancien_statut == nouveau_statut:
        return None

    return HistoriqueSousProjet.objects.create(
        sous_projet=sous_projet,
        numero_reception=(
            sous_projet.numero_reception_formulaire or ''
        ),
        intitule=(
            sous_projet.intitule_sous_projet or ''
        ),
        paysage=(
            sous_projet.paysage.nom
            if sous_projet.paysage
            else ''
        ),
        ancien_statut=ancien_statut,
        nouveau_statut=nouveau_statut,
        username=utilisateur.username or '',
    )


def filtrer_et_trier_preselection(request, sous_projets_base):
    """
    Applique les filtres communs aux pages de présélection :
    - Paysage / ZOCA
    - Numéro de réception
    - Tri numérique croissant du numéro de réception
    """
    selected_paysage_id = request.GET.get('paysage') or ''
    numero_reception = (request.GET.get('numero_reception') or '').strip()

    paysages_disponibles = (
        sous_projets_base
        .exclude(paysage__isnull=True)
        .values('paysage_id', 'paysage__nom', 'wilaya__nom')
        .distinct()
        .order_by('wilaya__nom', 'paysage__nom')
    )

    sous_projets = sous_projets_base
    selected_paysage = None

    if selected_paysage_id:
        sous_projets = sous_projets.filter(paysage_id=selected_paysage_id)
        try:
            selected_paysage = Paysage.objects.select_related(
                'commune__moughataa__wilaya'
            ).get(pk=selected_paysage_id)
        except Paysage.DoesNotExist:
            selected_paysage = None

    if numero_reception:
        sous_projets = sous_projets.filter(
            numero_reception_formulaire__icontains=numero_reception
        )

    sous_projets = (
        sous_projets
        .select_related('wilaya', 'moughataa', 'commune', 'paysage')
        .annotate(
            numero_reception_numerique=Case(
                When(
                    numero_reception_formulaire__regex=r'^[0-9]+$',
                    then=Cast('numero_reception_formulaire', IntegerField())
                ),
                default=Value(None),
                output_field=IntegerField()
            )
        )
        .order_by(
            'numero_reception_numerique',
            'numero_reception_formulaire',
            'id'
        )
    )

    return {
        'sous_projets': sous_projets,
        'paysages_disponibles': paysages_disponibles,
        'selected_paysage_id': str(selected_paysage_id),
        'selected_paysage': selected_paysage,
        'numero_reception': numero_reception,
    }


@login_required
def preselection_automatique(request):
    """
    Présélection automatique avec filtres, tri numérique,
    colonne Paysage / ZOCA et statistiques dynamiques.
    """
    utilisateur = get_current_user(request)
    sous_projets_base = get_accessible_sous_projets(utilisateur)

    filtre = filtrer_et_trier_preselection(request, sous_projets_base)
    sous_projets = filtre['sous_projets']

    lignes = []
    nb_rejetes = 0
    nb_a_examiner = 0

    for sp in sous_projets:
        evaluation = evaluer_preselection(sp)

        lignes.append({
            **evaluation,
            'paysage': sp.paysage.nom if sp.paysage else 'Non renseigné',
            'wilaya': sp.wilaya.nom if sp.wilaya else 'Non renseignée',
        })

        if evaluation['decision'] == 'Rejeté':
            nb_rejetes += 1
        else:
            nb_a_examiner += 1

    can_view_preselection_detail = bool(
        utilisateur
        and utilisateur.role in ['admin', 'superadmin', 'prescomite']
    )

    context = {
        'lignes': lignes,
        'total_dossiers': len(lignes),
        'nb_rejetes': nb_rejetes,
        'nb_a_examiner': nb_a_examiner,
        'paysages_disponibles': filtre['paysages_disponibles'],
        'selected_paysage_id': filtre['selected_paysage_id'],
        'selected_paysage': filtre['selected_paysage'],
        'numero_reception': filtre['numero_reception'],
        'can_view_preselection_detail': can_view_preselection_detail,
    }

    return render(
        request,
        'formulaire/preselection_automatique.html',
        context
    )

def evaluer_criteres_preselection(sp):
    """
    Évalue les critères de présélection un par un.
    Si le critère n'est pas mesurable automatiquement avec les données disponibles,
    on met 'À examiner'.
    """
    criteres = []

    def badge_for(decision):
        if decision == "Éligible":
            return "badge-ok"
        if decision == "Non éligible":
            return "badge-no"
        return "badge-review"

    # 1. Zone d’intervention
    if sp.wilaya and sp.commune and sp.paysage:
        criteres.append({
            "numero": 1,
            "critere": "Zone d’intervention",
            "question": "Le sous-projet est-il situé dans une zone couverte par le PADISAM/FCP ?",
            "resultat": "Oui",
            "decision": "Éligible",
            "motif": "Le sous-projet est localisé dans une zone renseignée du projet.",
            "badge_class": badge_for("Éligible"),
        })
    else:
        criteres.append({
            "numero": 1,
            "critere": "Zone d’intervention",
            "question": "Le sous-projet est-il situé dans une zone couverte par le PADISAM/FCP ?",
            "resultat": "À examiner",
            "decision": "À examiner",
            "motif": "La localisation n’est pas suffisamment renseignée pour confirmer automatiquement la zone.",
            "badge_class": badge_for("À examiner"),
        })

    # 2. Guichet concerné
    if sp.guichet in ["AGR", "ACI"]:
        criteres.append({
            "numero": 2,
            "critere": "Guichet concerné",
            "question": "Le dossier relève-t-il du bon guichet : AGR ou ACI ?",
            "resultat": sp.guichet,
            "decision": "Éligible",
            "motif": "Le guichet est renseigné dans le dossier.",
            "badge_class": badge_for("Éligible"),
        })
    else:
        criteres.append({
            "numero": 2,
            "critere": "Guichet concerné",
            "question": "Le dossier relève-t-il du bon guichet : AGR ou ACI ?",
            "resultat": "À examiner",
            "decision": "À examiner",
            "motif": "Le guichet n’est pas clairement exploitable automatiquement.",
            "badge_class": badge_for("À examiner"),
        })

    # 3. Profil du promoteur
    if sp.nom_statut_juridique:
        criteres.append({
            "numero": 3,
            "critere": "Profil du promoteur",
            "question": "Le demandeur fait-il partie des catégories admises ?",
            "resultat": "À examiner",
            "decision": "À examiner",
            "motif": "Le profil exact du promoteur nécessite une lecture métier du statut juridique.",
            "badge_class": badge_for("À examiner"),
        })
    else:
        criteres.append({
            "numero": 3,
            "critere": "Profil du promoteur",
            "question": "Le demandeur fait-il partie des catégories admises ?",
            "resultat": "Non renseigné",
            "decision": "À examiner",
            "motif": "Le profil du promoteur n’est pas suffisamment renseigné.",
            "badge_class": badge_for("À examiner"),
        })

    # 4. Statut du promoteur
    if sp.nom_statut_juridique:
        criteres.append({
            "numero": 4,
            "critere": "Statut du promoteur",
            "question": "Le statut du promoteur est-il conforme aux critères d’éligibilité ?",
            "resultat": "À examiner",
            "decision": "À examiner",
            "motif": "Le statut juridique est renseigné mais nécessite une validation métier.",
            "badge_class": badge_for("À examiner"),
        })
    else:
        criteres.append({
            "numero": 4,
            "critere": "Statut du promoteur",
            "question": "Le statut du promoteur est-il conforme aux critères d’éligibilité ?",
            "resultat": "Non renseigné",
            "decision": "À examiner",
            "motif": "Le statut du promoteur n’est pas renseigné.",
            "badge_class": badge_for("À examiner"),
        })

    # 5. Origine du promoteur
    criteres.append({
        "numero": 5,
        "critere": "Origine du promoteur",
        "question": "Le promoteur est-il résident ou originaire de la zone concernée ?",
        "resultat": "À examiner",
        "decision": "À examiner",
        "motif": "Cette information n’est pas disponible de manière fiable dans les champs actuels.",
        "badge_class": badge_for("À examiner"),
    })

    # 6. Dossier complet
    champs_min = [
        sp.numero_reception_formulaire,
        sp.intitule_sous_projet,
        sp.guichet,
        sp.type_projet,
        sp.nom_statut_juridique,
        sp.personne_contact_nom,
        sp.telephone,
        sp.wilaya,
        sp.moughataa,
        sp.commune,
        sp.paysage,
        sp.village,
        sp.objectif_sous_projet,
    ]
    if all(champs_min):
        criteres.append({
            "numero": 6,
            "critere": "Dossier complet",
            "question": "Le formulaire F1 et les pièces exigées sont-ils présents ?",
            "resultat": "Partiellement mesurable",
            "decision": "Éligible",
            "motif": "Les principaux champs du formulaire sont renseignés. Les pièces jointes restent à vérifier.",
            "badge_class": badge_for("Éligible"),
        })
    else:
        criteres.append({
            "numero": 6,
            "critere": "Dossier complet",
            "question": "Le formulaire F1 et les pièces exigées sont-ils présents ?",
            "resultat": "Incomplet",
            "decision": "Non éligible",
            "motif": "Le formulaire contient des informations essentielles manquantes.",
            "badge_class": badge_for("Non éligible"),
        })

    # 7. Activité proposée
    if sp.activites.exists():
        criteres.append({
            "numero": 7,
            "critere": "Activité proposée",
            "question": "L’activité figure-t-elle parmi les activités finançables du FCP ?",
            "resultat": "À examiner",
            "decision": "À examiner",
            "motif": "Les activités sont renseignées mais leur admissibilité doit être vérifiée selon la grille métier.",
            "badge_class": badge_for("À examiner"),
        })
    else:
        criteres.append({
            "numero": 7,
            "critere": "Activité proposée",
            "question": "L’activité figure-t-elle parmi les activités finançables du FCP ?",
            "resultat": "Aucune activité",
            "decision": "Non éligible",
            "motif": "Aucune activité n’est renseignée dans le dossier.",
            "badge_class": badge_for("Non éligible"),
        })

    # 8. Filière / chaîne de valeur
    if sp.chaine_approvisionnement:
        criteres.append({
            "numero": 8,
            "critere": "Filière / chaîne de valeur",
            "question": "Le sous-projet s’inscrit-il dans une filière admissible et prioritaire ?",
            "resultat": "À examiner",
            "decision": "À examiner",
            "motif": "La chaîne de valeur est renseignée mais doit être confrontée aux priorités de zone.",
            "badge_class": badge_for("À examiner"),
        })
    else:
        criteres.append({
            "numero": 8,
            "critere": "Filière / chaîne de valeur",
            "question": "Le sous-projet s’inscrit-il dans une filière admissible et prioritaire ?",
            "resultat": "Non renseigné",
            "decision": "À examiner",
            "motif": "La chaîne de valeur n’est pas renseignée.",
            "badge_class": badge_for("À examiner"),
        })

    # 9. Marché potentiel
    if sp.marches_vises:
        criteres.append({
            "numero": 9,
            "critere": "Marché potentiel",
            "question": "Le marché potentiel est-il identifié et accessible ?",
            "resultat": "Oui, déclaré",
            "decision": "Éligible",
            "motif": "Le dossier identifie un marché ou des débouchés visés.",
            "badge_class": badge_for("Éligible"),
        })
    else:
        criteres.append({
            "numero": 9,
            "critere": "Marché potentiel",
            "question": "Le marché potentiel est-il identifié et accessible ?",
            "resultat": "Non",
            "decision": "À examiner",
            "motif": "Aucun marché visé n’est renseigné.",
            "badge_class": badge_for("À examiner"),
        })

    # 10. Cohérence du sous-projet
    if sp.segment_ca and sp.activites.exists():
        criteres.append({
            "numero": 10,
            "critere": "Cohérence du sous-projet",
            "question": "Le segment de la chaîne de valeur correspond-il au type de sous-projet demandé ?",
            "resultat": "À examiner",
            "decision": "À examiner",
            "motif": "Des données existent mais la cohérence nécessite une lecture métier.",
            "badge_class": badge_for("À examiner"),
        })
    else:
        criteres.append({
            "numero": 10,
            "critere": "Cohérence du sous-projet",
            "question": "Le segment de la chaîne de valeur correspond-il au type de sous-projet demandé ?",
            "resultat": "Insuffisant",
            "decision": "À examiner",
            "motif": "Le segment ou les activités ne permettent pas une validation automatique.",
            "badge_class": badge_for("À examiner"),
        })

    # 11. Coûts réalistes
    grand_total = (
        (sp.infrastructures.aggregate(total=Sum('montant_total'))['total'] or 0) +
        (sp.equipements.aggregate(total=Sum('montant_total'))['total'] or 0) +
        (sp.intrants.aggregate(total=Sum('montant_total'))['total'] or 0) +
        (sp.fonctionnements.aggregate(total=Sum('montant_total'))['total'] or 0) +
        (sp.services.aggregate(total=Sum('montant_total'))['total'] or 0)
    )

    if grand_total > 0:
        criteres.append({
            "numero": 11,
            "critere": "Coûts réalistes",
            "question": "Les coûts annoncés sont-ils cohérents avec les barèmes du FCP ?",
            "resultat": f"Total déclaré : {grand_total}",
            "decision": "À examiner",
            "motif": "Les montants existent mais leur conformité aux barèmes doit être vérifiée.",
            "badge_class": badge_for("À examiner"),
        })
    else:
        criteres.append({
            "numero": 11,
            "critere": "Coûts réalistes",
            "question": "Les coûts annoncés sont-ils cohérents avec les barèmes du FCP ?",
            "resultat": "Aucun coût exploitable",
            "decision": "À examiner",
            "motif": "Aucun coût consolidé n’est disponible pour une évaluation automatique.",
            "badge_class": badge_for("À examiner"),
        })

    # 12. Plafond de subvention
    total_subvention = (
        (sp.infrastructures.aggregate(total=Sum('subvention_padisam'))['total'] or 0) +
        (sp.equipements.aggregate(total=Sum('subvention_padisam'))['total'] or 0) +
        (sp.intrants.aggregate(total=Sum('subvention_padisam'))['total'] or 0) +
        (sp.services.aggregate(total=Sum('subvention_padisam'))['total'] or 0)
    )

    criteres.append({
        "numero": 12,
        "critere": "Plafond de subvention",
        "question": "La subvention demandée respecte-t-elle le plafond prévu ?",
        "resultat": f"Subvention déclarée : {total_subvention}",
        "decision": "À examiner",
        "motif": "Le montant de subvention doit être comparé aux plafonds par catégorie de promoteur.",
        "badge_class": badge_for("À examiner"),
    })

    # 13. Expérience du promoteur
    if sp.annee_debut_activites:
        criteres.append({
            "numero": 13,
            "critere": "Expérience du promoteur",
            "question": "Le promoteur a-t-il l’ancienneté ou l’expérience minimale requise ?",
            "resultat": f"Début activité : {sp.annee_debut_activites}",
            "decision": "Éligible",
            "motif": "Une ancienneté est renseignée dans le dossier.",
            "badge_class": badge_for("Éligible"),
        })
    else:
        criteres.append({
            "numero": 13,
            "critere": "Expérience du promoteur",
            "question": "Le promoteur a-t-il l’ancienneté ou l’expérience minimale requise ?",
            "resultat": "Non renseigné",
            "decision": "À examiner",
            "motif": "L’ancienneté ou l’expérience n’est pas renseignée.",
            "badge_class": badge_for("À examiner"),
        })

    # 14. Réalisme de l’extension
    if sp.realisations.exists():
        criteres.append({
            "numero": 14,
            "critere": "Réalisme de l’extension",
            "question": "L’augmentation d’activité annoncée est-elle raisonnable ?",
            "resultat": "Historique disponible",
            "decision": "À examiner",
            "motif": "Les réalisations passées existent mais le réalisme de l’extension doit être analysé.",
            "badge_class": badge_for("À examiner"),
        })
    else:
        criteres.append({
            "numero": 14,
            "critere": "Réalisme de l’extension",
            "question": "L’augmentation d’activité annoncée est-elle raisonnable ?",
            "resultat": "Pas d’historique",
            "decision": "À examiner",
            "motif": "Aucun historique suffisant pour juger automatiquement l’extension.",
            "badge_class": badge_for("À examiner"),
        })

    # 15. Moyens minimums disponibles
    if sp.ressources_promoteur:
        criteres.append({
            "numero": 15,
            "critere": "Moyens minimums disponibles",
            "question": "Le promoteur dispose-t-il du site, du minimum d’équipement et des ressources techniques de base ?",
            "resultat": "Ressources déclarées",
            "decision": "Éligible",
            "motif": "Des ressources du promoteur sont renseignées dans le dossier.",
            "badge_class": badge_for("Éligible"),
        })
    else:
        criteres.append({
            "numero": 15,
            "critere": "Moyens minimums disponibles",
            "question": "Le promoteur dispose-t-il du site, du minimum d’équipement et des ressources techniques de base ?",
            "resultat": "Non renseigné",
            "decision": "À examiner",
            "motif": "Les moyens minimums ne sont pas décrits dans le dossier.",
            "badge_class": badge_for("À examiner"),
        })

    # 16. Eau / faisabilité du site
    if sp.type_projet == 'AG':
        if sp.ressources_promoteur:
            criteres.append({
                "numero": 16,
                "critere": "Eau / faisabilité du site",
                "question": "Pour un sous-projet agricole, la source d’eau existe-t-elle ?",
                "resultat": "À examiner",
                "decision": "À examiner",
                "motif": "Les ressources sont déclarées mais la disponibilité réelle de l’eau doit être confirmée.",
                "badge_class": badge_for("À examiner"),
            })
        else:
            criteres.append({
                "numero": 16,
                "critere": "Eau / faisabilité du site",
                "question": "Pour un sous-projet agricole, la source d’eau existe-t-elle ?",
                "resultat": "Non renseigné",
                "decision": "À examiner",
                "motif": "Aucune information exploitable sur l’eau n’est disponible.",
                "badge_class": badge_for("À examiner"),
            })
    else:
        criteres.append({
            "numero": 16,
            "critere": "Eau / faisabilité du site",
            "question": "Pour un sous-projet agricole, la source d’eau existe-t-elle ?",
            "resultat": "Non applicable",
            "decision": "Éligible",
            "motif": "Ce critère ne s’applique pas directement à ce type de projet.",
            "badge_class": badge_for("Éligible"),
        })

    # 17. Double financement / duplication
    criteres.append({
        "numero": 17,
        "critere": "Double financement / duplication",
        "question": "Le sous-projet finance-t-il les mêmes dépenses qu’un autre appui déjà obtenu ?",
        "resultat": "À examiner",
        "decision": "À examiner",
        "motif": "Aucune information suffisante dans la base actuelle pour détecter automatiquement une duplication.",
        "badge_class": badge_for("À examiner"),
    })

    # 18. Exigences environnementales et sociales
    criteres.append({
        "numero": 18,
        "critere": "Exigences environnementales et sociales",
        "question": "Le sous-projet respecte-t-il les exigences sociales et environnementales de base ?",
        "resultat": "À examiner",
        "decision": "À examiner",
        "motif": "Le screening E&S n’est pas automatisable avec les champs actuels.",
        "badge_class": badge_for("À examiner"),
    })

    # 19. Antécédents de crédit
    total_rembourse = sum(emp.montant_rembourse or 0 for emp in sp.emprunts.all())
    total_emprunte = sum(emp.montant_emprunte or 0 for emp in sp.emprunts.all())

    if total_emprunte == 0:
        criteres.append({
            "numero": 19,
            "critere": "Antécédents de crédit",
            "question": "Le demandeur a-t-il des antécédents connus de non-remboursement ?",
            "resultat": "Aucun emprunt déclaré",
            "decision": "Éligible",
            "motif": "Aucun antécédent d’emprunt bloquant n’est déclaré.",
            "badge_class": badge_for("Éligible"),
        })
    elif total_rembourse <= total_emprunte:
        criteres.append({
            "numero": 19,
            "critere": "Antécédents de crédit",
            "question": "Le demandeur a-t-il des antécédents connus de non-remboursement ?",
            "resultat": "Données de crédit disponibles",
            "decision": "À examiner",
            "motif": "Les emprunts existent mais l’historique réel de remboursement doit être vérifié qualitativement.",
            "badge_class": badge_for("À examiner"),
        })
    else:
        criteres.append({
            "numero": 19,
            "critere": "Antécédents de crédit",
            "question": "Le demandeur a-t-il des antécédents connus de non-remboursement ?",
            "resultat": "Incohérent",
            "decision": "À examiner",
            "motif": "Les montants d’emprunt et de remboursement sont incohérents et nécessitent vérification.",
            "badge_class": badge_for("À examiner"),
        })

    # 20. Visite de terrain nécessaire
    criteres.append({
        "numero": 20,
        "critere": "Visite de terrain nécessaire",
        "question": "Les informations du dossier suffisent-elles ou une vérification sur site est-elle nécessaire ?",
        "resultat": "Visite recommandée",
        "decision": "À examiner",
        "motif": "Plusieurs critères nécessitent une validation physique ou documentaire complémentaire.",
        "badge_class": badge_for("À examiner"),
    })

    return criteres


@login_required
def preselection_detail(request, pk):
    utilisateur = get_current_user(request)

    if not utilisateur or utilisateur.role not in ['admin', 'superadmin', 'prescomite']:
        messages.error(request, "❌ Vous n'êtes pas autorisé à consulter le détail de la présélection.")
        return redirect('formulaire:preselection_automatique')

    sous_projet = get_object_or_404(get_accessible_sous_projets(utilisateur), pk=pk)

    criteres = evaluer_criteres_preselection(sous_projet)

    nb_eligibles = sum(1 for c in criteres if c['decision'] == 'Éligible')
    nb_non_eligibles = sum(1 for c in criteres if c['decision'] == 'Non éligible')
    nb_a_examiner = sum(1 for c in criteres if c['decision'] == 'À examiner')

    return render(request, 'formulaire/preselection_detail.html', {
        'sous_projet': sous_projet,
        'criteres': criteres,
        'nb_eligibles': nb_eligibles,
        'nb_non_eligibles': nb_non_eligibles,
        'nb_a_examiner': nb_a_examiner,
    })

@login_required
def preselection_comite_liste(request):
    """
    Liste des dossiers pour la présélection comité avec :
    - droits d'accès selon le rôle ;
    - filtre Paysage / ZOCA ;
    - recherche par N° de réception ;
    - tri numérique croissant ;
    - statistiques dynamiques.
    """
    utilisateur = get_current_user(request)

    if not can_view_comite_list(utilisateur):
        messages.error(
            request,
            "❌ Vous n'êtes pas autorisé à accéder à cette page."
        )
        return redirect('formulaire:accueil')

    sous_projets_base = get_accessible_sous_projets(utilisateur)
    filtre = filtrer_et_trier_preselection(request, sous_projets_base)
    sous_projets = filtre['sous_projets']

    lignes = []
    nb_preselectionnes = 0
    nb_rejetes = 0
    nb_a_examiner = 0
    nb_reexaminer = 0

    for sp in sous_projets:
        auto = appliquer_preselection_automatique(sp)
        decision_comite = sp.decision_comite or 'a_examiner'

        if decision_comite == 'preselectionne':
          nb_preselectionnes += 1
        elif decision_comite == 'rejete':
          nb_rejetes += 1
        elif decision_comite == 'reexaminer':
         nb_reexaminer += 1
        else:
         nb_a_examiner += 1

        lignes.append({
            'sous_projet': sp,
            'numero_reception': (
                sp.numero_reception_formulaire or 'Non renseigné'
            ),
            'paysage': (
                sp.paysage.nom if sp.paysage else 'Non renseigné'
            ),
            'wilaya': (
                sp.wilaya.nom if sp.wilaya else 'Non renseignée'
            ),
            'intitule': (
                sp.intitule_sous_projet or 'Non renseigné'
            ),
            'decision_automatique': auto['decision'],
            'motif_automatique': auto['motif'],
            'badge_class': auto['badge_class'],
            'bloque_comite': auto['bloque_comite'],
            'decision_comite': decision_comite,
            'motif_comite': sp.motif_comite,
            'status': sp.status,
            'score_comite': getattr(sp, 'score_comite', 0),
            'evaluation_comite_terminee': getattr(
                sp,
                'evaluation_comite_terminee',
                False
            ),
        })

    context = {
        'lignes': lignes,
        'total_dossiers': len(lignes),
        'nb_preselectionnes': nb_preselectionnes,
        'nb_rejetes': nb_rejetes,
        'nb_a_examiner': nb_a_examiner,
        'nb_reexaminer': nb_reexaminer,
        'paysages_disponibles': filtre['paysages_disponibles'],
        'selected_paysage_id': filtre['selected_paysage_id'],
        'selected_paysage': filtre['selected_paysage'],
        'numero_reception': filtre['numero_reception'],
        'current_list_url': request.get_full_path(),
        'can_validate_comite': can_validate_comite(utilisateur),
        'can_view_comite_list': True,
        'is_superadmin': is_superadmin(utilisateur),
    }

    return render(
        request,
        'formulaire/preselection_comite_liste.html',
        context
    )

@login_required
def preselection_comite_detail(request, pk):
    utilisateur = get_current_user(request)

    if not can_validate_comite(utilisateur):
        messages.error(
            request,
            "❌ Vous n'êtes pas autorisé à accéder au détail de la présélection comité."
        )
        return redirect('formulaire:liste_sous_projets')

    liste_retour = request.GET.get('next') or request.POST.get('next')

    if not liste_retour or not url_has_allowed_host_and_scheme(
        url=liste_retour,
        allowed_hosts={request.get_host()},
        require_https=request.is_secure(),
    ):
        liste_retour = reverse(
            'formulaire:preselection_comite_liste'
        )

    sous_projet = get_object_or_404(
        get_accessible_sous_projets(utilisateur),
        pk=pk
    )

    auto = appliquer_preselection_automatique(sous_projet)

    # Aucun verrouillage : le comité peut toujours modifier la décision.
    readonly_mode = False

    if not sous_projet.resultats_preselection.exists():
        criteres_auto = evaluer_criteres_preselection(
            sous_projet
        )

        for item in criteres_auto:
            decision_auto = item.get(
                'decision',
                'À examiner'
            )

            if decision_auto == 'Éligible':
                decision_auto_db = 'eligible'
            elif decision_auto == 'Non éligible':
                decision_auto_db = 'non_eligible'
            else:
                decision_auto_db = 'a_examiner'

            ResultatPreselection.objects.create(
                sous_projet=sous_projet,
                numero_critere=item.get('numero'),
                critere=item.get('critere'),
                question=item.get('question'),
                resultat_automatique=item.get('resultat'),
                decision_automatique=decision_auto_db,
                motif_automatique=item.get('motif'),
            )

    queryset = (
        sous_projet
        .resultats_preselection
        .all()
        .order_by('numero_critere')
    )

    if request.method == 'POST':
        ancien_statut = normaliser_statut_preselection(
            sous_projet.decision_comite
            or sous_projet.status
            or 'a_examiner'
        )

        formset = ResultatPreselectionFormSet(
            request.POST,
            queryset=queryset,
            prefix='critere'
        )

        decision_form = DecisionComiteSousProjetForm(
            request.POST,
            instance=sous_projet,
            prefix='global'
        )

        decision_globale = request.POST.get(
            'global-decision_comite'
        )

        motif_global = request.POST.get(
            'global-motif_comite'
        )

        # Une décision complète impose la saisie des décisions et motifs,
        # mais ne verrouille plus le dossier.
        decision_finale = decision_globale in [
            'preselectionne',
            'rejete',
        ]

        # Décisions provisoires : le dossier reste modifiable.
        decision_attente = decision_globale in [
            '',
            None,
            'reexaminer',
        ]

        if decision_finale:
            for form in formset.forms:
                decision = form.data.get(
                    form.add_prefix(
                        'decision_comite'
                    )
                )

                motif = form.data.get(
                    form.add_prefix(
                        'motif_comite'
                    )
                )

                if not decision:
                    form.add_error(
                        'decision_comite',
                        (
                            "La décision est obligatoire "
                            "pour valider définitivement "
                            "le dossier."
                        )
                    )

                if not motif:
                    form.add_error(
                        'motif_comite',
                        (
                            "Le motif est obligatoire "
                            "pour valider définitivement "
                            "le dossier."
                        )
                    )

            if not motif_global:
                decision_form.add_error(
                    'motif_comite',
                    (
                        "Le motif global est obligatoire "
                        "pour valider définitivement "
                        "le dossier."
                    )
                )

        if formset.is_valid() and decision_form.is_valid():
            formset.save()

            sous_projet = decision_form.save(
                commit=False
            )

            if decision_finale:
                score = 0

                for resultat in (
                    sous_projet
                    .resultats_preselection
                    .all()
                ):
                    if (
                        resultat.decision_comite
                        == 'preselectionne'
                    ):
                        score += 1

                if hasattr(
                    sous_projet,
                    'score_comite'
                ):
                    sous_projet.score_comite = score

                if hasattr(
                    sous_projet,
                    'evaluation_comite_terminee'
                ):
                    sous_projet.evaluation_comite_terminee = True

                if hasattr(
                    sous_projet,
                    'date_evaluation_comite'
                ):
                    sous_projet.date_evaluation_comite = (
                        timezone.now()
                    )

                # Traçabilité du président ayant
                # enregistré la décision finale.
                sous_projet.president_comite = utilisateur

                sous_projet.president_username = (
                    utilisateur.username
                )

                sous_projet.president_nom_complet = (
                    f"{utilisateur.prenom or ''} "
                    f"{utilisateur.nom or ''}"
                ).strip() or utilisateur.username

                sous_projet.date_decision_comite = (
                    timezone.now()
                )

                # Le statut reflète la décision, sans bloquer les modifications futures.
                sous_projet.status = decision_globale

                sous_projet.save()

                enregistrer_historique_statut(
                    sous_projet=sous_projet,
                    utilisateur=utilisateur,
                    ancien_statut=ancien_statut,
                    nouveau_statut=decision_globale
                )

                messages.success(
                    request,
                    (
                        "✅ Décision du comité enregistrée avec succès. "
                        "Elle reste modifiable ultérieurement. "
                        f"Score : {score}/20."
                    )
                )

                return redirect(liste_retour)

            if decision_attente:
                # Le statut reste non final pour permettre
                # la modification du dossier.
                sous_projet.status = 'etude'

                if hasattr(
                    sous_projet,
                    'evaluation_comite_terminee'
                ):
                    sous_projet.evaluation_comite_terminee = False

                sous_projet.save()

                nouveau_statut = normaliser_statut_preselection(
                    decision_globale or 'a_examiner'
                )

                enregistrer_historique_statut(
                    sous_projet=sous_projet,
                    utilisateur=utilisateur,
                    ancien_statut=ancien_statut,
                    nouveau_statut=nouveau_statut
                )

                if decision_globale == 'reexaminer':
                    messages.success(
                        request,
                        (
                            "✅ Décision enregistrée : "
                            "le dossier est à réexaminer "
                            "et reste modifiable."
                        )
                    )

                else:
                    messages.success(
                        request,
                        (
                            "✅ Travail du comité enregistré. "
                            "Le dossier reste modifiable."
                        )
                    )

                return redirect(liste_retour)

        messages.error(
            request,
            "❌ Veuillez corriger les erreurs du formulaire."
        )

    else:
        formset = ResultatPreselectionFormSet(
            queryset=queryset,
            prefix='critere'
        )

        decision_form = DecisionComiteSousProjetForm(
            instance=sous_projet,
            prefix='global',
            readonly=False
        )

    criteres = list(queryset)

    nb_eligibles = sum(
        1
        for critere in criteres
        if critere.decision_automatique == 'eligible'
    )

    nb_non_eligibles = sum(
        1
        for critere in criteres
        if critere.decision_automatique == 'non_eligible'
    )

    nb_a_examiner = sum(
        1
        for critere in criteres
        if critere.decision_automatique == 'a_examiner'
    )

    context = {
        'sous_projet': sous_projet,
        'formset': formset,
        'decision_form': decision_form,
        'criteres': criteres,
        'nb_eligibles': nb_eligibles,
        'nb_non_eligibles': nb_non_eligibles,
        'nb_a_examiner': nb_a_examiner,
        'auto': auto,
        'readonly_mode': False,
        'liste_retour': liste_retour,
    }

    return render(
        request,
        'formulaire/preselection_comite_detail.html',
        context
    )

@login_required
def rapport_paysage_type_financement(request):
    """
    Rapport dynamique :
    - Filtre exclusif par Wilaya OU par Paysage / ZOCA
    - Répartition par type de projet : AG, EL, SER, ENV
    - Financement détaillé : Subvention, Contribution, Autre, Total
    - Total financement = somme réelle des montant_total
    """
    utilisateur = get_current_user(request)

    sous_projets_base = get_accessible_sous_projets(utilisateur)

    # =====================================================
    # Choix du type de filtre : wilaya OU paysage
    # =====================================================
    filter_by = request.GET.get('filter_by') or 'wilaya'

    selected_wilaya_id = ''
    selected_paysage_id = ''

    if filter_by == 'wilaya':
        selected_wilaya_id = request.GET.get('wilaya') or ''
    elif filter_by == 'paysage':
        selected_paysage_id = request.GET.get('paysage') or ''
    else:
        filter_by = 'wilaya'

    selected_wilaya = None
    selected_paysage = None

    # =====================================================
    # Listes disponibles pour les filtres
    # =====================================================
    wilayas_disponibles = (
        sous_projets_base
        .exclude(wilaya__isnull=True)
        .values('wilaya_id', 'wilaya__nom')
        .distinct()
        .order_by('wilaya__nom')
    )

    paysages_disponibles = (
        sous_projets_base
        .exclude(paysage__isnull=True)
        .values('paysage_id', 'paysage__nom', 'wilaya__nom')
        .distinct()
        .order_by('wilaya__nom', 'paysage__nom')
    )

    # =====================================================
    # Application du filtre choisi
    # =====================================================
    sous_projets = sous_projets_base

    if filter_by == 'wilaya' and selected_wilaya_id:
        sous_projets = sous_projets.filter(wilaya_id=selected_wilaya_id)

        try:
            selected_wilaya = Wilaya.objects.get(id=selected_wilaya_id)
        except Wilaya.DoesNotExist:
            selected_wilaya = None

    if filter_by == 'paysage' and selected_paysage_id:
        sous_projets = sous_projets.filter(paysage_id=selected_paysage_id)

        try:
            selected_paysage = Paysage.objects.get(id=selected_paysage_id)
        except Paysage.DoesNotExist:
            selected_paysage = None

    type_codes = ['AG', 'EL', 'SER', 'ENV']

    def total_relation(queryset, relation, champ):
        return queryset.aggregate(
            total=Sum(f'{relation}__{champ}')
        )['total'] or 0

    rapport = []

    wilayas = (
        sous_projets
        .exclude(wilaya__isnull=True)
        .values('wilaya_id', 'wilaya__nom')
        .distinct()
        .order_by('wilaya__nom')
    )

    for wilaya in wilayas:
        wilaya_id = wilaya['wilaya_id']
        wilaya_nom = wilaya['wilaya__nom'] or "Non renseignée"

        qs_wilaya = sous_projets.filter(wilaya_id=wilaya_id)

        paysages = (
            qs_wilaya
            .values('paysage_id', 'paysage__nom')
            .distinct()
            .order_by('paysage__nom')
        )

        lignes_paysages = []

        total_wilaya_types = {code: 0 for code in type_codes}
        total_wilaya_subvention = 0
        total_wilaya_contribution = 0
        total_wilaya_autre = 0
        total_wilaya_financement = 0
        total_wilaya_projets = 0

        for paysage in paysages:
            paysage_id = paysage['paysage_id']
            paysage_nom = paysage['paysage__nom'] or "Non renseigné"

            if paysage_id:
                qs_paysage = qs_wilaya.filter(paysage_id=paysage_id)
            else:
                qs_paysage = qs_wilaya.filter(paysage__isnull=True)

            total_projets_paysage = qs_paysage.count()

            types_count = {code: 0 for code in type_codes}

            stats_types = (
                qs_paysage
                .values('type_projet')
                .annotate(total=Count('id'))
            )

            for item in stats_types:
                type_code = item['type_projet']
                total_type = item['total']

                if type_code in type_codes:
                    types_count[type_code] = total_type
                    total_wilaya_types[type_code] += total_type

            total_financement = (
                total_relation(qs_paysage, 'infrastructures', 'montant_total') +
                total_relation(qs_paysage, 'equipements', 'montant_total') +
                total_relation(qs_paysage, 'intrants', 'montant_total') +
                total_relation(qs_paysage, 'fonctionnements', 'montant_total') +
                total_relation(qs_paysage, 'services', 'montant_total')
            )

            subvention = (
                total_relation(qs_paysage, 'infrastructures', 'subvention_padisam') +
                total_relation(qs_paysage, 'equipements', 'subvention_padisam') +
                total_relation(qs_paysage, 'intrants', 'subvention_padisam') +
                total_relation(qs_paysage, 'services', 'subvention_padisam')
            )

            contribution = (
                total_relation(qs_paysage, 'infrastructures', 'contribution_promoteur') +
                total_relation(qs_paysage, 'equipements', 'contribution_promoteur') +
                total_relation(qs_paysage, 'intrants', 'contribution_promoteur') +
                total_relation(qs_paysage, 'fonctionnements', 'contribution_promoteur') +
                total_relation(qs_paysage, 'services', 'contribution_promoteur')
            )

            autre = (
                total_relation(qs_paysage, 'infrastructures', 'autre_financement') +
                total_relation(qs_paysage, 'equipements', 'autre_financement') +
                total_relation(qs_paysage, 'intrants', 'autre_financement') +
                total_relation(qs_paysage, 'fonctionnements', 'autre_financement') +
                total_relation(qs_paysage, 'services', 'autre_financement')
            )

            total_wilaya_subvention += subvention
            total_wilaya_contribution += contribution
            total_wilaya_autre += autre
            total_wilaya_financement += total_financement
            total_wilaya_projets += total_projets_paysage

            lignes_paysages.append({
                'paysage': paysage_nom,
                'types': types_count,
                'total_projets': total_projets_paysage,
                'subvention': subvention,
                'contribution': contribution,
                'autre': autre,
                'total_financement': total_financement,
            })

        rapport.append({
            'wilaya': wilaya_nom,
            'paysages': lignes_paysages,
            'total_types': total_wilaya_types,
            'total_projets': total_wilaya_projets,
            'total_subvention': total_wilaya_subvention,
            'total_contribution': total_wilaya_contribution,
            'total_autre': total_wilaya_autre,
            'total_financement': total_wilaya_financement,
        })

    total_general_projets = sum(item['total_projets'] for item in rapport)
    total_general_subvention = sum(item['total_subvention'] for item in rapport)
    total_general_contribution = sum(item['total_contribution'] for item in rapport)
    total_general_autre = sum(item['total_autre'] for item in rapport)
    total_general_financement = sum(item['total_financement'] for item in rapport)

    total_general_types = {
        'AG': sum(item['total_types']['AG'] for item in rapport),
        'EL': sum(item['total_types']['EL'] for item in rapport),
        'SER': sum(item['total_types']['SER'] for item in rapport),
        'ENV': sum(item['total_types']['ENV'] for item in rapport),
    }

    context = {
        'filter_by': filter_by,

        'rapport': rapport,

        'selected_wilaya_id': str(selected_wilaya_id),
        'selected_paysage_id': str(selected_paysage_id),

        'selected_wilaya': selected_wilaya,
        'selected_paysage': selected_paysage,

        'wilayas_disponibles': wilayas_disponibles,
        'paysages_disponibles': paysages_disponibles,

        'total_general_types': total_general_types,
        'total_general_projets': total_general_projets,
        'total_general_subvention': total_general_subvention,
        'total_general_contribution': total_general_contribution,
        'total_general_autre': total_general_autre,
        'total_general_financement': total_general_financement,

        'user_name': request.session.get('user_name'),
        'user_role': request.session.get('user_role'),
        'user_wilaya_nom': request.session.get('user_wilaya_nom'),
    }

    return render(
        request,
        'formulaire/rapport_paysage_type_financement.html',
        context
    )


# =========================================================
# EXPORT EXCEL : RÉSULTATS DE PRÉSÉLECTION DU COMITÉ
# =========================================================

def exporter_rapport_preselection_excel(
    lignes,
    total_couts,
    total_subventions,
    selected_wilaya=None,
    selected_paysage=None
):
    """
    Génère le fichier Excel du rapport de présélection du comité.

    Colonnes exportées :
    - N° de réception
    - Intitulé
    - Type de projet
    - Guichet
    - Paysage / ZOCA
    - Village
    - Nom et statut juridique
    - Principal domaine d'activité
    - Personne contact
    - Téléphone
    - Coût total
    - Subvention
    - Critères rejetés automatiquement
    - Critères rejetés par le CPR
    - Décision
    - Motif
    """
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "Présélection comité"

    # =====================================================
    # COULEURS ET STYLES
    # =====================================================
    green_fill = PatternFill(
        "solid",
        fgColor="5DBB63"
    )

    light_green_fill = PatternFill(
        "solid",
        fgColor="E8F5E9"
    )

    gray_fill = PatternFill(
        "solid",
        fgColor="F3F4F6"
    )

    white_font = Font(
        color="FFFFFF",
        bold=True
    )

    bold_font = Font(
        bold=True
    )

    thin_gray = Side(
        style="thin",
        color="D1D5DB"
    )

    border = Border(
        left=thin_gray,
        right=thin_gray,
        top=thin_gray,
        bottom=thin_gray,
    )

    # =====================================================
    # TITRE
    # 16 colonnes : A à P
    # =====================================================
    worksheet.merge_cells("A1:P1")

    worksheet["A1"] = (
        "Résultats de présélection du comité"
    )

    worksheet["A1"].font = Font(
        bold=True,
        size=16,
        color="FFFFFF"
    )

    worksheet["A1"].fill = green_fill

    worksheet["A1"].alignment = Alignment(
        horizontal="center",
        vertical="center"
    )

    worksheet.row_dimensions[1].height = 26

    # =====================================================
    # SOUS-TITRE ET FILTRES
    # =====================================================
    wilaya_label = (
        selected_wilaya.nom
        if selected_wilaya
        else "Toutes les Wilayas"
    )

    paysage_label = (
        selected_paysage.nom
        if selected_paysage
        else "Tous les paysages"
    )

    worksheet.merge_cells("A2:P2")

    worksheet["A2"] = (
        f"Wilaya : {wilaya_label} | "
        f"Paysage / ZOCA : {paysage_label}"
    )

    worksheet["A2"].font = bold_font
    worksheet["A2"].fill = light_green_fill

    worksheet["A2"].alignment = Alignment(
        horizontal="center",
        vertical="center"
    )

    # =====================================================
    # EN-TÊTES
    # =====================================================
    headers = [
        "N°",
        "Intitulé du sous-projet",
        "Type de projet",
        "Guichet",
        "Paysage / ZOCA",
        "Village",
        "Nom et statut juridique",
        "Principal domaine d’activité",
        "Personne contact",
        "Téléphone",
        "Coût total",
        "Subvention",
        "Critères rejetés automatiquement",
        "Critères rejetés CPR",
        "Décision",
        "Motif",
    ]

    header_row = 4

    for column_index, header in enumerate(
        headers,
        start=1
    ):
        cell = worksheet.cell(
            row=header_row,
            column=column_index,
            value=header
        )

        cell.fill = green_fill
        cell.font = white_font

        cell.alignment = Alignment(
            horizontal="center",
            vertical="center",
            wrap_text=True
        )

        cell.border = border

    worksheet.row_dimensions[header_row].height = 42

    # Geler les lignes de titre et d'en-tête
    worksheet.freeze_panes = "A5"

    # Filtre Excel sur les 16 colonnes
    worksheet.auto_filter.ref = (
        f"A4:P{max(4, len(lignes) + 4)}"
    )

    # =====================================================
    # LIBELLÉS DES DÉCISIONS
    # =====================================================
    decision_labels = {
        "preselectionne": "Présélectionné",
        "rejete": "Rejeté",
        "reexaminer": "À réexaminer",
        "a_examiner": "À examiner",
        "": "À examiner",
        None: "À examiner",
    }

    # =====================================================
    # LIGNES DU RAPPORT
    # =====================================================
    for row_index, ligne in enumerate(
        lignes,
        start=5
    ):
        values = [
            ligne.get(
                "numero",
                "Non renseigné"
            ),

            ligne.get(
                "intitule",
                "Non renseigné"
            ),

            ligne.get(
                "type_projet",
                "Non renseigné"
            ),

            ligne.get(
                "guichet",
                "Non renseigné"
            ),

            ligne.get(
                "paysage",
                "Non renseigné"
            ),

            ligne.get(
                "village",
                "Non renseigné"
            ),

            ligne.get(
                "nom_statut_juridique",
                "Non renseigné"
            ),

            ligne.get(
                "domaine_activite",
                "Non renseigné"
            ),

            ligne.get(
                "personne_contact",
                "Non renseigné"
            ),

            ligne.get(
                "telephone",
                "Non renseigné"
            ),

            float(
                ligne.get("cout_total") or 0
            ),

            float(
                ligne.get("subvention") or 0
            ),

            ligne.get(
                "nb_rejets_automatiques",
                0
            ),

            ligne.get(
                "nb_rejets_comite",
                0
            ),

            decision_labels.get(
                ligne.get("decision_code"),
                ligne.get(
                    "decision",
                    "À examiner"
                )
            ),

            ligne.get(
                "motif",
                "Non renseigné"
            ),
        ]

        for column_index, value in enumerate(
            values,
            start=1
        ):
            cell = worksheet.cell(
                row=row_index,
                column=column_index,
                value=value
            )

            cell.border = border

            cell.alignment = Alignment(
                vertical="top",
                wrap_text=True
            )

            # Colonnes financières :
            # K = coût total
            # L = subvention
            if column_index in (11, 12):
                cell.number_format = '#,##0.00'
                cell.alignment = Alignment(
                    horizontal="right",
                    vertical="top"
                )

            # Colonnes centrées :
            # A : numéro
            # C : type
            # D : guichet
            # J : téléphone
            # M et N : nombres de critères
            # O : décision
            if column_index in (
                1,
                3,
                4,
                10,
                13,
                14,
                15,
            ):
                cell.alignment = Alignment(
                    horizontal="center",
                    vertical="top",
                    wrap_text=True
                )

    # =====================================================
    # LIGNE TOTAL GÉNÉRAL
    # =====================================================
    total_row = len(lignes) + 5

    # Fusion de A à J
    worksheet.merge_cells(
        start_row=total_row,
        start_column=1,
        end_row=total_row,
        end_column=10,
    )

    total_label_cell = worksheet.cell(
        row=total_row,
        column=1,
        value="TOTAL GÉNÉRAL"
    )

    total_label_cell.font = bold_font
    total_label_cell.fill = gray_fill

    total_label_cell.alignment = Alignment(
        horizontal="right",
        vertical="center"
    )

    total_label_cell.border = border

    # Bordures des cellules fusionnées A à J
    for merged_column in range(1, 11):
        worksheet.cell(
            row=total_row,
            column=merged_column
        ).border = border

        worksheet.cell(
            row=total_row,
            column=merged_column
        ).fill = gray_fill

    # Coût total en colonne K
    total_cost_cell = worksheet.cell(
        row=total_row,
        column=11,
        value=float(total_couts or 0)
    )

    total_cost_cell.font = bold_font
    total_cost_cell.fill = gray_fill
    total_cost_cell.number_format = '#,##0.00'
    total_cost_cell.border = border

    total_cost_cell.alignment = Alignment(
        horizontal="right",
        vertical="center"
    )

    # Subvention en colonne L
    total_subvention_cell = worksheet.cell(
        row=total_row,
        column=12,
        value=float(total_subventions or 0)
    )

    total_subvention_cell.font = bold_font
    total_subvention_cell.fill = gray_fill
    total_subvention_cell.number_format = '#,##0.00'
    total_subvention_cell.border = border

    total_subvention_cell.alignment = Alignment(
        horizontal="right",
        vertical="center"
    )

    # Colonnes M à P
    for column_index in range(13, 17):
        cell = worksheet.cell(
            row=total_row,
            column=column_index
        )

        cell.fill = gray_fill
        cell.border = border

    # =====================================================
    # LARGEURS DES COLONNES
    # =====================================================
    column_widths = {
        1: 14,   # N°
        2: 35,   # Intitulé
        3: 20,   # Type de projet
        4: 14,   # Guichet
        5: 26,   # Paysage / ZOCA
        6: 20,   # Village
        7: 32,   # Nom et statut juridique
        8: 34,   # Domaine d'activité
        9: 26,   # Personne contact
        10: 17,  # Téléphone
        11: 18,  # Coût total
        12: 18,  # Subvention
        13: 20,  # Rejets automatiques
        14: 18,  # Rejets CPR
        15: 20,  # Décision
        16: 42,  # Motif
    }

    for column_index, width in (
        column_widths.items()
    ):
        worksheet.column_dimensions[
            get_column_letter(column_index)
        ].width = width

    # Masquer le quadrillage Excel
    worksheet.sheet_view.showGridLines = False

    # =====================================================
    # RÉPONSE HTTP
    # =====================================================
    response = HttpResponse(
        content_type=(
            "application/vnd.openxmlformats-"
            "officedocument.spreadsheetml.sheet"
        )
    )

    response["Content-Disposition"] = (
        'attachment; '
        'filename="rapport_preselection_comite.xlsx"'
    )

    workbook.save(response)

    return response


@login_required
def rapport_preselection_comite(request):
    """
    Rapport détaillé des résultats de présélection du comité.

    Filtres :
    - Wilaya
    - Paysage / ZOCA
    - Type de projet
    - Décision du comité
    - Numéro de réception

    Une décision vide ou NULL est affichée comme « À examiner ».
    L'ancienne valeur « a_examiner » reste prise en charge pour
    la compatibilité avec les données existantes.
    """
    utilisateur = get_current_user(request)

    if not can_view_comite_list(utilisateur):
        messages.error(
            request,
            "❌ Vous n'êtes pas autorisé à consulter ce rapport."
        )
        return redirect('formulaire:accueil')

    sous_projets_base = get_accessible_sous_projets(utilisateur)

    selected_wilaya_id = (request.GET.get('wilaya') or '').strip()
    selected_paysage_id = (request.GET.get('paysage') or '').strip()
    selected_type_projet = (request.GET.get('type_projet') or '').strip()
    selected_decision = (request.GET.get('decision') or '').strip()
    numero_reception = (
        request.GET.get('numero_reception') or ''
    ).strip()

    wilayas_disponibles = (
        sous_projets_base
        .exclude(wilaya__isnull=True)
        .values('wilaya_id', 'wilaya__nom')
        .distinct()
        .order_by('wilaya__nom')
    )

    paysages_base = sous_projets_base
    if selected_wilaya_id:
        paysages_base = paysages_base.filter(
            wilaya_id=selected_wilaya_id
        )

    paysages_disponibles = (
        paysages_base
        .exclude(paysage__isnull=True)
        .values(
            'paysage_id',
            'paysage__nom',
            'wilaya__nom',
        )
        .distinct()
        .order_by(
            'wilaya__nom',
            'paysage__nom',
        )
    )

    sous_projets = sous_projets_base

    if selected_wilaya_id:
        sous_projets = sous_projets.filter(
            wilaya_id=selected_wilaya_id
        )

    if selected_paysage_id:
        sous_projets = sous_projets.filter(
            paysage_id=selected_paysage_id
        )

    if selected_type_projet:
        sous_projets = sous_projets.filter(
            type_projet=selected_type_projet
        )

    if numero_reception:
        sous_projets = sous_projets.filter(
            numero_reception_formulaire__icontains=numero_reception
        )

    if selected_decision == 'a_examiner':
        sous_projets = sous_projets.filter(
            Q(decision_comite__isnull=True)
            | Q(decision_comite='')
            | Q(decision_comite='a_examiner')
        )
    elif selected_decision:
        sous_projets = sous_projets.filter(
            decision_comite=selected_decision
        )

    sous_projets = (
        sous_projets
        .select_related(
            'wilaya',
            'moughataa',
            'commune',
            'paysage',
        )
        .prefetch_related(
            'infrastructures',
            'equipements',
            'intrants',
            'fonctionnements',
            'services',
            'resultats_preselection',
        )
        .annotate(
            numero_reception_numerique=Case(
                When(
                    numero_reception_formulaire__regex=r'^[0-9]+$',
                    then=Cast(
                        'numero_reception_formulaire',
                        IntegerField()
                    )
                ),
                default=Value(None),
                output_field=IntegerField()
            )
        )
        .order_by(
            'numero_reception_numerique',
            'numero_reception_formulaire',
            'id',
        )
    )

    lignes = []

    nb_preselectionnes = 0
    nb_rejetes = 0
    nb_reexaminer = 0
    nb_a_examiner = 0

    total_couts = 0
    total_subventions = 0
    total_criteres_rejetes_automatiques = 0
    total_criteres_rejetes_comite = 0

    decision_labels = {
        'preselectionne': 'Présélectionné',
        'rejete': 'Rejeté',
        'reexaminer': 'À réexaminer',
        'a_examiner': 'À examiner',
        '': 'À examiner',
        None: 'À examiner',
    }

    for sp in sous_projets:
        decision_code = sp.decision_comite or 'a_examiner'

        if decision_code == 'preselectionne':
            nb_preselectionnes += 1
        elif decision_code == 'rejete':
            nb_rejetes += 1
        elif decision_code == 'reexaminer':
            nb_reexaminer += 1
        else:
            nb_a_examiner += 1

        infra_totaux = compute_financement_totals(
            sp.infrastructures.all(),
            include_subvention=True
        )
        equip_totaux = compute_financement_totals(
            sp.equipements.all(),
            include_subvention=True
        )
        intrant_totaux = compute_financement_totals(
            sp.intrants.all(),
            include_subvention=True
        )
        fonctionnement_totaux = compute_financement_totals(
            sp.fonctionnements.all(),
            include_subvention=False
        )
        service_totaux = compute_financement_totals(
            sp.services.all(),
            include_subvention=True
        )

        cout_total = (
            infra_totaux['total_montant']
            + equip_totaux['total_montant']
            + intrant_totaux['total_montant']
            + fonctionnement_totaux['total_montant']
            + service_totaux['total_montant']
        )

        total_subvention = (
            infra_totaux['total_subvention']
            + equip_totaux['total_subvention']
            + intrant_totaux['total_subvention']
            + service_totaux['total_subvention']
        )

        resultats = list(sp.resultats_preselection.all())

        nb_rejets_automatiques = sum(
            1
            for resultat in resultats
            if resultat.decision_automatique == 'non_eligible'
        )

        nb_rejets_comite = sum(
            1
            for resultat in resultats
            if resultat.decision_comite == 'rejete'
        )

        total_couts += cout_total
        total_subventions += total_subvention
        total_criteres_rejetes_automatiques += nb_rejets_automatiques
        total_criteres_rejetes_comite += nb_rejets_comite

        # Afficher uniquement la valeur réellement saisie dans
        # « Principal domaine d'activités ».
        # Ne jamais remplacer une valeur vide par le type de projet
        # (AG, EL, SER ou ENV).
        domaine_activite = (
            (sp.principal_domaine_activites or '').strip()
            or 'Non renseigné'
        )

        lignes.append({
            'sous_projet': sp,
            'numero': sp.numero_reception_formulaire or '-',
            'intitule': sp.intitule_sous_projet or '-',
            'type_projet': (
    sp.get_type_projet_display()
    if sp.type_projet
    else 'Non renseigné'
),

'guichet': (
    sp.get_guichet_display()
    if sp.guichet
    else 'Non renseigné'
),
'paysage': (
    sp.paysage.nom
    if sp.paysage
    else 'Non renseigné'
),
            'village': sp.village or '-',
            'nom_statut_juridique': sp.nom_statut_juridique or '-',
            'domaine_activite': domaine_activite,
            'personne_contact': sp.personne_contact_nom or '-',
            'telephone': sp.telephone or '-',
            'cout_total': cout_total,
            'subvention': total_subvention,
            'nb_rejets_automatiques': nb_rejets_automatiques,
            'nb_rejets_comite': nb_rejets_comite,
            'decision_code': decision_code,
            'decision': decision_labels.get(
                decision_code,
                'À examiner'
            ),
            'motif': sp.motif_comite or '-',
        })

    total_dossiers = len(lignes)

    nb_dossiers_traites = (
        nb_preselectionnes
        + nb_rejetes
        + nb_reexaminer
    )

    taux_preselection = (
        round(
            nb_preselectionnes * 100 / total_dossiers,
            2
        )
        if total_dossiers
        else 0
    )

    taux_rejet = (
        round(
            nb_rejetes * 100 / total_dossiers,
            2
        )
        if total_dossiers
        else 0
    )

    taux_traitement = (
        round(
            nb_dossiers_traites * 100 / total_dossiers,
            2
        )
        if total_dossiers
        else 0
    )

    chart_decision_labels = [
        'Présélectionnés',
        'Rejetés',
        'À réexaminer',
        'À examiner',
    ]
    chart_decision_data = [
        nb_preselectionnes,
        nb_rejetes,
        nb_reexaminer,
        nb_a_examiner,
    ]

    selected_wilaya = None
    selected_paysage = None

    if selected_wilaya_id:
        try:
            selected_wilaya = Wilaya.objects.get(
                pk=selected_wilaya_id
            )
        except (Wilaya.DoesNotExist, ValueError):
            selected_wilaya = None

    if selected_paysage_id:
        try:
            selected_paysage = Paysage.objects.get(
                pk=selected_paysage_id
            )
        except (Paysage.DoesNotExist, ValueError):
            selected_paysage = None

    # Export Excel avec exactement les filtres actuellement appliqués.
    if request.GET.get('export') == 'excel':
        return exporter_rapport_preselection_excel(
            lignes=lignes,
            total_couts=total_couts,
            total_subventions=total_subventions,
            selected_wilaya=selected_wilaya,
            selected_paysage=selected_paysage,
        )

    type_projet_choices = [
        ('AG', 'Agriculture'),
        ('EL', 'Élevage'),
        ('SER', 'Service'),
        ('ENV', 'Environnement'),
    ]

    context = {
        'lignes': lignes,

        'total_dossiers': total_dossiers,
        'nb_preselectionnes': nb_preselectionnes,
        'nb_rejetes': nb_rejetes,
        'nb_reexaminer': nb_reexaminer,
        'nb_a_examiner': nb_a_examiner,
        'nb_dossiers_traites': nb_dossiers_traites,

        'taux_preselection': taux_preselection,
        'taux_rejet': taux_rejet,
        'taux_traitement': taux_traitement,

        'total_couts': total_couts,
        'total_subventions': total_subventions,
        'total_criteres_rejetes_automatiques': (
            total_criteres_rejetes_automatiques
        ),
        'total_criteres_rejetes_comite': (
            total_criteres_rejetes_comite
        ),

        'chart_decision_labels': json.dumps(
            chart_decision_labels
        ),
        'chart_decision_data': json.dumps(
            chart_decision_data
        ),

        'wilayas_disponibles': wilayas_disponibles,
        'paysages_disponibles': paysages_disponibles,
        'type_projet_choices': type_projet_choices,

        'selected_wilaya_id': str(selected_wilaya_id),
        'selected_paysage_id': str(selected_paysage_id),
        'selected_type_projet': selected_type_projet,
        'selected_decision': selected_decision,
        'numero_reception': numero_reception,

        'selected_wilaya': selected_wilaya,
        'selected_paysage': selected_paysage,

        'user_name': request.session.get('user_name'),
        'user_role': request.session.get('user_role'),
        'user_wilaya_nom': request.session.get('user_wilaya_nom'),
    }

    return render(
        request,
        'formulaire/rapport_preselection_comite.html',
        context
    )

def enregistrer_derniere_modification(sous_projet, utilisateur):
    """Enregistre le dernier modificateur et actualise date_modification."""
    if not sous_projet or not utilisateur:
        return

    sous_projet.dernier_modificateur_username = utilisateur.username
    sous_projet.save(update_fields=[
        'dernier_modificateur_username',
        'date_modification',
    ])


def enregistrer_historique_modification(
    sous_projet,
    utilisateur,
    section,
    action='modification'
):
    """
    Enregistre une action réalisée sur un dossier :
    création, modification d'une section ou suppression.

    Les informations principales du dossier sont copiées afin que
    l'historique reste lisible même si le dossier est supprimé plus tard.
    """
    if not sous_projet or not utilisateur:
        return None

    return HistoriqueModificationDossier.objects.create(
        sous_projet=sous_projet,
        numero_reception=(
            sous_projet.numero_reception_formulaire or ''
        ),
        intitule=(
            sous_projet.intitule_sous_projet or ''
        ),
        paysage=(
            sous_projet.paysage.nom
            if sous_projet.paysage
            else ''
        ),
        action=action,
        section=section or '',
        username=utilisateur.username or '',
    )


@login_required
def modifier_section_sous_projet(request, pk, section):
    """
    Modification d'une section depuis la page de détail.

    Pour les formsets, on utilise toujours extra=0 :
    - seules les lignes déjà enregistrées sont affichées ;
    - une nouvelle ligne apparaît uniquement après un clic sur
      « + Ajouter une ligne » ;
    - TOTAL_FORMS indique à Django le nombre réel de lignes envoyées.

    Le management_form du formset doit rester présent dans le template.
    """

    utilisateur = get_current_user(request)

    sous_projet = get_object_or_404(
        get_accessible_sous_projets(utilisateur),
        pk=pk
    )

    if not can_modify_project(utilisateur, sous_projet):
        messages.error(
            request,
            "❌ Vous n'êtes pas autorisé à modifier ce sous-projet."
        )
        return redirect(
            'formulaire:detail_sous_projet',
            pk=sous_projet.pk
        )

    section_config = {
        'informations': {
            'titre': 'Modifier les informations générales',
            'type': 'form',
            'form_class': SousProjetForm,
            'prefix': None,
        },
        'activites': {
            'titre': 'Modifier les activités',
            'type': 'inline_formset',
            'formset_class': None,
            'prefix': 'activite',
        },
        'infrastructures': {
            'titre': 'Modifier les infrastructures',
            'type': 'inline_formset',
            'formset_class': InfrastructureFormSet,
            'prefix': 'infra',
        },
        'equipements': {
            'titre': 'Modifier les équipements',
            'type': 'inline_formset',
            'formset_class': EquipementFormSet,
            'prefix': 'equip',
        },
        'intrants': {
            'titre': 'Modifier les intrants',
            'type': 'inline_formset',
            'formset_class': IntrantFormSet,
            'prefix': 'intrant',
        },
        'fonctionnements': {
            'titre': 'Modifier le fonctionnement',
            'type': 'inline_formset',
            'formset_class': FonctionnementFormSet,
            'prefix': 'fonc',
        },
        'services': {
            'titre': 'Modifier les services',
            'type': 'inline_formset',
            'formset_class': ServiceFormSet,
            'prefix': 'serv',
        },
        'realisations': {
            'titre': 'Modifier les réalisations passées du promoteur',
            'type': 'manual_formset',
            'formset_class': RealisationFormSet,
            'prefix': 'real',
        },
        'emprunts': {
            'titre': 'Modifier les passifs / emprunts',
            'type': 'manual_formset',
            'formset_class': EmpruntFormSet,
            'prefix': 'emprunt',
        },
        'promoteur': {
            'titre': 'Modifier les renseignements promoteur',
            'type': 'form',
            'form_class': PromoteurFinalForm,
            'prefix': None,
        },
    }

    if section not in section_config:
        messages.error(request, "❌ Section introuvable.")
        return redirect(
            'formulaire:detail_sous_projet',
            pk=sous_projet.pk
        )

    config = section_config[section]
    form = None
    formset = None

    # Aucune ligne vide n'est affichée automatiquement.
    # Une nouvelle ligne est créée uniquement lorsque l'utilisateur
    # clique sur le bouton « + Ajouter une ligne » dans le template.
    # En POST, le nombre réel de lignes est fourni par TOTAL_FORMS.
    extra_value = 0

    if config['type'] == 'inline_formset':
        if section == 'activites':
            DynamicFormSet = get_activite_formset_class(
                extra=extra_value
            )
        else:
            DynamicFormSet = clone_formset_with_extra(
                config['formset_class'],
                extra_value
            )

    elif config['type'] == 'manual_formset':
        DynamicFormSet = clone_formset_with_extra(
            config['formset_class'],
            extra_value
        )

    if request.method == 'POST':

        if config['type'] == 'form':
            if section == 'informations':
                form = config['form_class'](
                    request.POST,
                    instance=sous_projet,
                    user=utilisateur
                )
            else:
                form = config['form_class'](
                    request.POST,
                    instance=sous_projet
                )

            if form.is_valid():
                obj = form.save(commit=False)

                if section == 'informations':
                    if (
                        is_agent_saisie(utilisateur)
                        and utilisateur.wilaya_id
                    ):
                        obj.wilaya = utilisateur.wilaya

                    # Le créateur initial du dossier ne doit pas changer.
                    obj.createur_username = (
                        sous_projet.createur_username
                    )

                try:
                    with transaction.atomic():
                        obj.dernier_modificateur_username = (
                            utilisateur.username
                        )
                        obj.save()

                        enregistrer_historique_modification(
                            sous_projet=obj,
                            utilisateur=utilisateur,
                            section=section,
                            action='modification'
                        )

                except Exception as exc:
                    messages.error(
                        request,
                        "❌ La modification a été annulée : "
                        f"{exc}"
                    )
                    return render(
                        request,
                        'formulaire/modifier_section_sous_projet.html',
                        {
                            'sous_projet': sous_projet,
                            'section': section,
                            'titre_section': config['titre'],
                            'form': form,
                            'formset': None,
                            'config_type': 'form',
                            'can_add_line': False,
                            'formset_prefix': '',
                            'user_name': request.session.get('user_name'),
                            'user_role': request.session.get('user_role'),
                            'user_wilaya_nom': request.session.get(
                                'user_wilaya_nom'
                            ),
                        }
                    )

                messages.success(
                    request,
                    "✅ Section modifiée avec succès."
                )
                return redirect(
                    'formulaire:detail_sous_projet',
                    pk=sous_projet.pk
                )

        elif config['type'] == 'manual_formset':
            formset = DynamicFormSet(
                request.POST,
                prefix=config['prefix']
            )

            if formset.is_valid():
                try:
                    with transaction.atomic():

                        if section == 'realisations':
                            sous_projet.realisations.all().delete()

                            for current_form in formset.forms:
                                data = (
                                    getattr(
                                        current_form,
                                        'cleaned_data',
                                        None
                                    )
                                    or {}
                                )

                                if data.get('DELETE'):
                                    continue

                                has_data = any([
                                    data.get('produit'),
                                    data.get('annee_1'),
                                    data.get('volume_annee_1'),
                                    data.get('ventes_usd_annee_1'),
                                    data.get('prix_vente_mru_annee_1'),
                                    data.get('annee_2'),
                                    data.get('volume_annee_2'),
                                    data.get('ventes_usd_annee_2'),
                                    data.get('prix_vente_mru_annee_2'),
                                    data.get('annee_3'),
                                    data.get('volume_annee_3'),
                                    data.get('ventes_usd_annee_3'),
                                    data.get('prix_vente_mru_annee_3'),
                                ])

                                if not has_data:
                                    continue

                                RealisationPassee.objects.create(
                                    sous_projet=sous_projet,
                                    produit=data.get('produit'),
                                    annee_1=data.get('annee_1'),
                                    volume_annee_1=data.get(
                                        'volume_annee_1'
                                    ),
                                    ventes_usd_annee_1=data.get(
                                        'ventes_usd_annee_1'
                                    ),
                                    prix_vente_mru_annee_1=data.get(
                                        'prix_vente_mru_annee_1'
                                    ),
                                    annee_2=data.get('annee_2'),
                                    volume_annee_2=data.get(
                                        'volume_annee_2'
                                    ),
                                    ventes_usd_annee_2=data.get(
                                        'ventes_usd_annee_2'
                                    ),
                                    prix_vente_mru_annee_2=data.get(
                                        'prix_vente_mru_annee_2'
                                    ),
                                    annee_3=data.get('annee_3'),
                                    volume_annee_3=data.get(
                                        'volume_annee_3'
                                    ),
                                    ventes_usd_annee_3=data.get(
                                        'ventes_usd_annee_3'
                                    ),
                                    prix_vente_mru_annee_3=data.get(
                                        'prix_vente_mru_annee_3'
                                    ),
                                )

                        elif section == 'emprunts':
                            sous_projet.emprunts.all().delete()

                            for current_form in formset.forms:
                                data = (
                                    getattr(
                                        current_form,
                                        'cleaned_data',
                                        None
                                    )
                                    or {}
                                )

                                if data.get('DELETE'):
                                    continue

                                has_data = any([
                                    data.get('annee'),
                                    data.get(
                                        'institution_financiere'
                                    ),
                                    data.get('montant_emprunte'),
                                    data.get('montant_rembourse'),
                                ])

                                if not has_data:
                                    continue

                                PassifEmprunt.objects.create(
                                    sous_projet=sous_projet,
                                    annee=data.get('annee'),
                                    institution_financiere=data.get(
                                        'institution_financiere'
                                    ),
                                    montant_emprunte=data.get(
                                        'montant_emprunte'
                                    ),
                                    montant_rembourse=data.get(
                                        'montant_rembourse'
                                    ),
                                )

                        enregistrer_derniere_modification(
                            sous_projet,
                            utilisateur
                        )

                        enregistrer_historique_modification(
                            sous_projet=sous_projet,
                            utilisateur=utilisateur,
                            section=section,
                            action='modification'
                        )

                    messages.success(
                        request,
                        "✅ Section modifiée avec succès."
                    )
                    return redirect(
                        'formulaire:detail_sous_projet',
                        pk=sous_projet.pk
                    )

                except Exception as exc:
                    messages.error(
                        request,
                        "❌ La modification a été annulée : "
                        f"{exc}"
                    )

        elif config['type'] == 'inline_formset':
            formset = DynamicFormSet(
                request.POST,
                instance=sous_projet,
                prefix=config['prefix']
            )

            if formset.is_valid():
                try:
                    with transaction.atomic():
                        formset.save()

                        enregistrer_derniere_modification(
                            sous_projet,
                            utilisateur
                        )

                        enregistrer_historique_modification(
                            sous_projet=sous_projet,
                            utilisateur=utilisateur,
                            section=section,
                            action='modification'
                        )

                    messages.success(
                        request,
                        "✅ Section modifiée avec succès."
                    )
                    return redirect(
                        'formulaire:detail_sous_projet',
                        pk=sous_projet.pk
                    )

                except Exception as exc:
                    messages.error(
                        request,
                        "❌ La modification a été annulée : "
                        f"{exc}"
                    )

    else:
        if config['type'] == 'form':
            if section == 'informations':
                form = config['form_class'](
                    instance=sous_projet,
                    user=utilisateur
                )
            else:
                form = config['form_class'](
                    instance=sous_projet
                )

        elif config['type'] == 'inline_formset':
            formset = DynamicFormSet(
                instance=sous_projet,
                prefix=config['prefix']
            )

        elif config['type'] == 'manual_formset':
            if section == 'realisations':
                initial = list(
                    sous_projet.realisations.values(
                        'produit',
                        'annee_1',
                        'volume_annee_1',
                        'ventes_usd_annee_1',
                        'prix_vente_mru_annee_1',
                        'annee_2',
                        'volume_annee_2',
                        'ventes_usd_annee_2',
                        'prix_vente_mru_annee_2',
                        'annee_3',
                        'volume_annee_3',
                        'ventes_usd_annee_3',
                        'prix_vente_mru_annee_3',
                    )
                )
            else:
                initial = list(
                    sous_projet.emprunts.values(
                        'annee',
                        'institution_financiere',
                        'montant_emprunte',
                        'montant_rembourse',
                    )
                )

            formset = DynamicFormSet(
                initial=initial,
                prefix=config['prefix']
            )

    return render(
        request,
        'formulaire/modifier_section_sous_projet.html',
        {
            'sous_projet': sous_projet,
            'section': section,
            'titre_section': config['titre'],
            'form': form,
            'formset': formset,
            'config_type': (
                'formset'
                if config['type'] in [
                    'inline_formset',
                    'manual_formset',
                ]
                else config['type']
            ),
            # Informations utilisées par le JavaScript du template
            # pour cloner formset.empty_form et augmenter TOTAL_FORMS.
            'can_add_line': config['type'] in [
                'inline_formset',
                'manual_formset',
            ],
            'formset_prefix': config.get('prefix') or '',
            'user_name': request.session.get('user_name'),
            'user_role': request.session.get('user_role'),
            'user_wilaya_nom': request.session.get(
                'user_wilaya_nom'
            ),
        }
    )
def superadmin_required(view_func):
    """
    Autorise uniquement les utilisateurs ayant le rôle superadmin.
    """
    @wraps(view_func)
    def wrapper(request, *args, **kwargs):
        utilisateur = get_current_user(request)

        if not utilisateur:
            return redirect('formulaire:login')

        if utilisateur.role != 'superadmin':
            messages.error(
                request,
                "❌ Cette page est réservée au Super Administrateur."
            )
            return redirect('formulaire:accueil')

        return view_func(request, *args, **kwargs)

    return wrapper


@superadmin_required
def admin_historique_modifications(request):
    """
    Historique des créations, modifications et suppressions de dossiers.
    """

    historiques = (
        HistoriqueModificationDossier.objects
        .select_related('sous_projet')
        .all()
    )

    numero_reception = (
        request.GET.get('numero_reception') or ''
    ).strip()

    username = (
        request.GET.get('username') or ''
    ).strip()

    action = (
        request.GET.get('action') or ''
    ).strip()

    section = (
        request.GET.get('section') or ''
    ).strip()

    date_debut = (
        request.GET.get('date_debut') or ''
    ).strip()

    date_fin = (
        request.GET.get('date_fin') or ''
    ).strip()

    if numero_reception:
        historiques = historiques.filter(
            numero_reception__icontains=numero_reception
        )

    if username:
        historiques = historiques.filter(
            username__icontains=username
        )

    if action:
        historiques = historiques.filter(
            action=action
        )

    if section:
        historiques = historiques.filter(
            section=section
        )

    if date_debut:
        historiques = historiques.filter(
            date_action__date__gte=date_debut
        )

    if date_fin:
        historiques = historiques.filter(
            date_action__date__lte=date_fin
        )

    historiques = historiques.order_by(
        '-date_action',
        '-id'
    )

    paginator = Paginator(historiques, 30)
    page_number = request.GET.get('page')
    page_obj = paginator.get_page(page_number)

    sections_disponibles = (
        HistoriqueModificationDossier.objects
        .exclude(section='')
        .values_list('section', flat=True)
        .distinct()
        .order_by('section')
    )

    context = {
        'page_obj': page_obj,
        'historiques': page_obj.object_list,

        'numero_reception': numero_reception,
        'username': username,
        'selected_action': action,
        'selected_section': section,
        'date_debut': date_debut,
        'date_fin': date_fin,

        'sections_disponibles': sections_disponibles,

        'total_resultats': paginator.count,

        'user_name': request.session.get('user_name'),
        'user_role': request.session.get('user_role'),
        'user_wilaya_nom': request.session.get(
            'user_wilaya_nom'
        ),
    }

    return render(
        request,
        'formulaire/admin_historique_modifications.html',
        context
    )


@superadmin_required
def admin_historique_preselection(request):
    """
    Historique des changements de décision du comité.
    """

    historiques = (
        HistoriqueSousProjet.objects
        .select_related('sous_projet')
        .all()
    )

    numero_reception = (
        request.GET.get('numero_reception') or ''
    ).strip()

    username = (
        request.GET.get('username') or ''
    ).strip()

    ancien_statut = (
        request.GET.get('ancien_statut') or ''
    ).strip()

    nouveau_statut = (
        request.GET.get('nouveau_statut') or ''
    ).strip()

    date_debut = (
        request.GET.get('date_debut') or ''
    ).strip()

    date_fin = (
        request.GET.get('date_fin') or ''
    ).strip()

    if numero_reception:
        historiques = historiques.filter(
            numero_reception__icontains=numero_reception
        )

    if username:
        historiques = historiques.filter(
            username__icontains=username
        )

    if ancien_statut:
        historiques = historiques.filter(
            ancien_statut=ancien_statut
        )

    if nouveau_statut:
        historiques = historiques.filter(
            nouveau_statut=nouveau_statut
        )

    if date_debut:
        historiques = historiques.filter(
            date_action__date__gte=date_debut
        )

    if date_fin:
        historiques = historiques.filter(
            date_action__date__lte=date_fin
        )

    historiques = historiques.order_by(
        '-date_action',
        '-id'
    )

    paginator = Paginator(historiques, 30)
    page_number = request.GET.get('page')
    page_obj = paginator.get_page(page_number)

    statut_choices = [
        ('a_examiner', 'À examiner'),
        ('preselectionne', 'Présélectionné'),
        ('rejete', 'Rejeté'),
        ('reexaminer', 'À réexaminer'),
    ]

    context = {
        'page_obj': page_obj,
        'historiques': page_obj.object_list,

        'numero_reception': numero_reception,
        'username': username,
        'selected_ancien_statut': ancien_statut,
        'selected_nouveau_statut': nouveau_statut,
        'date_debut': date_debut,
        'date_fin': date_fin,

        'statut_choices': statut_choices,

        'total_resultats': paginator.count,

        'user_name': request.session.get('user_name'),
        'user_role': request.session.get('user_role'),
        'user_wilaya_nom': request.session.get(
            'user_wilaya_nom'
        ),
    }

    return render(
        request,
        'formulaire/admin_historique_preselection.html',
        context
    )
